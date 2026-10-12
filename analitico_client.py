# -*- coding: utf-8 -*-
"""Leitura e indicadores da base compartilhada ANALITICO.xlsx."""

from __future__ import annotations

import math
import os
import re
import sys
import unicodedata
from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

from excel_client import DataClient, DataClientError, _localizar_excel_no_onedrive


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _key(value) -> str:
    value = unicodedata.normalize("NFKD", _text(value).upper())
    value = "".join(char for char in value if not unicodedata.combining(char))
    return re.sub(r"[^A-Z0-9]+", "_", value).strip("_")


def _date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = _text(value).replace("T", " ").split(" ", 1)[0]
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw, pattern).date()
        except ValueError:
            continue
    return None


def _hours(value) -> float | None:
    """Converte duração Excel, horas numéricas ou HH:MM[:SS] em horas."""
    if value in (None, ""):
        return None
    if isinstance(value, timedelta):
        return value.total_seconds() / 3600
    if isinstance(value, datetime):
        return value.hour + value.minute / 60 + value.second / 3600
    if hasattr(value, "hour") and hasattr(value, "minute"):
        return value.hour + value.minute / 60 + getattr(value, "second", 0) / 3600
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if isinstance(value, float) and math.isnan(value):
            return None
        # Durações do Excel são armazenadas como fração de um dia.
        return float(value) * 24 if abs(float(value)) < 1 else float(value)
    raw = _text(value).replace(",", ".")
    match = re.fullmatch(r"(-?\d+):([0-5]?\d)(?::([0-5]?\d))?", raw)
    if match:
        sign = -1 if match.group(1).startswith("-") else 1
        return sign * (abs(int(match.group(1))) + int(match.group(2)) / 60 + int(match.group(3) or 0) / 3600)
    try:
        return float(raw)
    except ValueError:
        return None


def _truthy(value) -> bool:
    normalized = _key(value)
    return bool(normalized and normalized not in {"NAO", "N", "0", "FALSO", "FALSE", "NA"})


class AnaliticoClient:
    """Localiza e consulta ANALITICO.xlsx sem alterar o arquivo."""

    _path_cache: str | None = None

    def __init__(self, excel_path=None):
        self._explicit_path = Path(excel_path).expanduser() if excel_path else None

    def _resolve_path(self) -> Path:
        if self._explicit_path:
            if self._explicit_path.is_file():
                return self._explicit_path
            raise DataClientError(f"O arquivo não foi encontrado: {self._explicit_path}", status_code=404)
        configured = os.getenv("ANALITICO_PATH", "").strip()
        if configured:
            path = Path(configured).expanduser()
            if path.is_file():
                return path
            raise DataClientError(f"ANALITICO_PATH não foi encontrado: {path}", status_code=404)
        filename = os.getenv("ANALITICO_FILENAME", "ANALITICO.xlsx").strip() or "ANALITICO.xlsx"
        if Path(filename).name != filename or Path(filename).suffix.casefold() != ".xlsx":
            raise DataClientError("ANALITICO_FILENAME deve conter somente o nome de um arquivo .xlsx.", status_code=500)
        cached = self.__class__._path_cache
        if cached and Path(cached).is_file():
            return Path(cached)
        directories = []
        try:
            directories.append(Path(DataClient()._resolver_caminho_excel()).parent)
        except DataClientError:
            pass
        directories.extend([
            Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent,
            Path.cwd(), Path.home() / "Desktop", Path.home() / "Área de Trabalho",
            Path.home() / "Documents", Path.home() / "Documentos",
        ])
        seen = set()
        for directory in directories:
            marker = str(directory).casefold()
            if marker in seen:
                continue
            seen.add(marker)
            candidate = directory / filename
            if candidate.is_file() and not candidate.name.startswith("~$"):
                self.__class__._path_cache = str(candidate)
                return candidate
        found, _roots = _localizar_excel_no_onedrive(filename)
        if found:
            self.__class__._path_cache = str(found)
            return found
        raise DataClientError(
            "O arquivo 'ANALITICO.xlsx' não foi encontrado ao lado da base principal, "
            "na Área de Trabalho, em Documentos ou no OneDrive sincronizado.", status_code=404
        )

    @staticmethod
    def _find_sheet(workbook):
        required = {_key("BD"), _key("ABERTURA"), _key("LP15")}
        for sheet in workbook.worksheets:
            headers = {_key(cell.value) for cell in sheet[1] if cell.value}
            if required.issubset(headers):
                return sheet
        raise DataClientError("ANALITICO.xlsx não possui uma aba com as colunas BD, ABERTURA e LP15.", status_code=422)

    def _read(self) -> list[dict]:
        path = self._resolve_path()
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except PermissionError as exc:
            raise DataClientError("ANALITICO.xlsx está bloqueado. Feche-o no Excel e tente novamente.", status_code=423) from exc
        try:
            sheet = self._find_sheet(workbook)
            headers = [_key(cell.value) for cell in sheet[1]]
            rows = []
            for values in sheet.iter_rows(min_row=2, values_only=True):
                if not any(value not in (None, "") for value in values):
                    continue
                rows.append({headers[index]: value for index, value in enumerate(values) if index < len(headers) and headers[index]})
            return rows
        finally:
            workbook.close()

    def status(self):
        path = self._resolve_path()
        return {"arquivo": path.name, "caminho": str(path), "modo": "Excel compartilhado (somente leitura)"}

    def options(self):
        rows = self._read()
        return {
            "tecnologias": sorted({_text(row.get("TECNOLOGIA")) for row in rows if _text(row.get("TECNOLOGIA"))}, key=str.casefold),
        }

    def dashboard(self, *, start=None, end=None, technologies=None):
        rows = self._read()
        start = _date(start)
        end = _date(end)
        selected = {_key(value) for value in (technologies or []) if _text(value)}
        if selected:
            rows = [row for row in rows if _key(row.get("TECNOLOGIA")) in selected]

        dated = [(row, _date(row.get("ABERTURA"))) for row in rows]
        period_rows = [row for row, opened in dated if opened and (not start or opened >= start) and (not end or opened <= end)]
        by_month = Counter(opened.strftime("%Y-%m") for _row, opened in dated if opened and (not start or opened >= start) and (not end or opened <= end))

        def grouped(field, limit=10, minimum=2):
            counts = Counter(_text(row.get(field)) for row in period_rows if _text(row.get(field)))
            return [
                {"nome": name, "quantidade": count}
                for name, count in counts.most_common(limit) if count >= minimum
            ]

        r30 = {}
        for row in period_rows:
            if not _truthy(row.get("R30_TRATATIVAS")):
                continue
            lp = _text(row.get("LP15")) or "Não informado"
            item = r30.setdefault(lp.casefold(), {"lp15": lp, "cliente": _text(row.get("CLIENTE")), "quantidade": 0})
            item["quantidade"] += 1
            item["cliente"] = _text(row.get("CLIENTE")) or item["cliente"]

        today = date.today()
        critical = []
        for days in (30, 15, 7):
            counts = Counter()
            for row, opened in dated:
                if not opened or opened < today - timedelta(days=days - 1) or opened > today:
                    continue
                if not _truthy(row.get("CLIENTE_CRITICO")):
                    continue
                name = _text(row.get("CLIENTE")) or _text(row.get("LP15")) or "Não informado"
                counts[name] += 1
            critical.append({"dias": days, "itens": [{"cliente": name, "quantidade": count} for name, count in counts.most_common(10)]})

        tmrs = []
        for row in period_rows:
            elapsed, paused = _hours(row.get("TEMPO_CORRIDO")), _hours(row.get("TEMPO_PARADA"))
            if elapsed is not None:
                tmrs.append(max(0, elapsed - (paused or 0)))

        months = sorted(by_month)
        return {
            "status": self.status(),
            "bds_por_mes": {"labels": months, "values": [by_month[key] for key in months]},
            "armarios_criticos": grouped("ARMARIO"),
            "erbs_criticas": grouped("ERB"),
            "r30_tratativas": sorted(r30.values(), key=lambda item: (-item["quantidade"], item["lp15"].casefold()))[:10],
            "clientes_criticos": critical,
            "tmr_medio_horas": round(sum(tmrs) / len(tmrs), 2) if tmrs else 0,
            "total_periodo": len(period_rows),
        }
