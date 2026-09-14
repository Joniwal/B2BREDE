# -*- coding: utf-8 -*-
"""Acesso ao arquivo ATIVACAO.xlsx usado pela página de Ativação.

O cliente preserva a tabela estruturada existente e aceita os nomes atuais
das colunas da planilha. Na interface, ``DATA VENCIMENTO`` é apresentada como
Data de Agendamento e ``DATA ENCERRAMENTO`` como Data de Execução.
"""

from __future__ import annotations

import os
import sys
import tempfile
import threading
import unicodedata
from collections import Counter
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils.cell import range_boundaries

from excel_client import DataClient, DataClientError, _localizar_excel_no_onedrive


ATIVACAO_LOCK = threading.RLock()

FIELDS = [
    "id", "cliente", "cidade", "servico", "tecnologia", "empresa",
    "status", "data_agendamento", "data_execucao", "no_mes", "tecnico",
    "faturado", "com_rfs",
]

EXCEL_HEADERS = [
    "ID", "CLIENTES", "CIDADE", "SERVIÇO", "TECNOLOGIA", "EMPRESA",
    "STATUS", "DATA VENCIMENTO", "DATA ENCERRAMENTO", "NO MÊS",
    "EXECUTADO POR", "FATURADO", "COM RFS",
]

HEADER_ALIASES = {
    "id": {"ID"},
    "cliente": {"CLIENTE", "CLIENTES"},
    "cidade": {"CIDADE"},
    "servico": {"SERVICO"},
    "tecnologia": {"TECNOLOGIA"},
    "empresa": {"EMPRESA"},
    "status": {"STATUS"},
    "data_agendamento": {"DATA AGENDAMENTO", "DATA VENCIMENTO"},
    "data_execucao": {"DATA EXECUCAO", "DATA ENCERRAMENTO"},
    "no_mes": {"NO MES"},
    "tecnico": {"TECNICO", "EXECUTADO POR"},
    "faturado": {"FATURADO"},
    "com_rfs": {"COM RFS", "RFS"},
}

FIXED_OPTIONS = {
    "tecnologias": ["ERB", "GPON"],
    "empresas": ["VIVO", "RS TELECOM"],
    "status": ["OK", "NOK"],
    "sim_nao": ["SIM", "NÃO"],
    "tecnicos": [
        "ALEXSANDRO NUNES DA SILVA",
        "ANDRE DOS SANTOS CARVALHO",
        "EDSON LUIS DE CRISTO",
        "EMERSON FAVERO RODRIGUES",
        "LUCAS SILVA ANDRADE",
        "MARCOS ROBERTO HOLTMAN",
        "ERENILSON SANT'ANA",
    ],
}

REQUIRED_LABELS = {
    "cliente": "Cliente",
    "cidade": "Cidade",
    "servico": "Serviço",
    "tecnologia": "Tecnologia",
    "empresa": "Empresa",
    "status": "Status",
    "data_agendamento": "Data de Agendamento",
    "no_mes": "No Mês",
    "tecnico": "Técnico",
    "faturado": "Faturado",
    "com_rfs": "Com RFS",
}


def _normalize_header(value) -> str:
    text = "" if value is None else str(value).strip().upper()
    text = "".join(
        char for char in unicodedata.normalize("NFKD", text)
        if not unicodedata.combining(char)
    )
    return " ".join(text.split())


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _parse_id(value, *, required=False) -> int | None:
    raw = _text(value)
    if not raw:
        if required:
            raise DataClientError("Campo ID obrigatório.", status_code=400)
        return None
    if not raw.isdigit() or int(raw) < 1:
        raise DataClientError(
            "ID inválido. Informe um número inteiro positivo.",
            status_code=400,
        )
    return int(raw)


def _parse_date(value, *, required=False, label="Data") -> date | None:
    if value is None or _text(value) == "":
        if required:
            raise DataClientError(f"Campo {label} obrigatório.", status_code=400)
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = _text(value)
    for pattern in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(raw[:10], pattern).date()
        except ValueError:
            continue
    raise DataClientError(f"{label} inválida.", status_code=400)


def _unique_sorted(values):
    unique = {_text(value) for value in values if _text(value)}
    return sorted(unique, key=str.casefold)


class AtivacaoClient:
    """CRUD e agregações da tabela ATIVACAO."""

    # Guarda o arquivo localizado automaticamente. Se ele for movido ou
    # renomeado durante a execução, a existência é conferida antes do reuso e
    # uma nova busca é realizada.
    _caminho_cache: str | None = None

    def __init__(self, excel_path: str | os.PathLike | None = None):
        self._explicit_path = Path(excel_path).expanduser() if excel_path else None

    @staticmethod
    def _candidate_directories(main_path: Path | None = None) -> list[Path]:
        """Pastas rápidas para uso portátil, antes da busca no OneDrive."""
        candidates: list[Path] = []
        if getattr(sys, "frozen", False):
            candidates.append(Path(sys.executable).resolve().parent)
        else:
            candidates.append(Path(__file__).resolve().parent)

        candidates.append(Path.cwd())
        home = Path.home()
        candidates.extend(
            (
                home / "Desktop",
                home / "Área de Trabalho",
                home / "Documents",
                home / "Documentos",
            )
        )
        # Mantém compatibilidade com a organização anterior, mas somente
        # depois dos locais portáteis. Isso evita que uma planilha vazia
        # criada por uma versão antiga esconda o arquivo compartilhado.
        if main_path:
            candidates.append(main_path.parent)

        result: list[Path] = []
        seen: set[str] = set()
        for directory in candidates:
            try:
                key = str(directory.resolve()).casefold()
            except OSError:
                key = str(directory).casefold()
            if key not in seen:
                seen.add(key)
                result.append(directory)
        return result

    def _resolve_path(self) -> Path:
        if self._explicit_path:
            return self._explicit_path
        configured = os.getenv("ATIVACAO_PATH", "").strip()
        if configured:
            return Path(configured).expanduser()

        filename = os.getenv("ATIVACAO_FILENAME", "ATIVACAO.xlsx").strip() or "ATIVACAO.xlsx"
        if Path(filename).name != filename or Path(filename).suffix.casefold() != ".xlsx":
            raise DataClientError(
                "ATIVACAO_FILENAME deve conter somente o nome de um arquivo .xlsx.",
                status_code=500,
            )

        cached = self.__class__._caminho_cache
        if cached:
            cached_path = Path(cached)
            if cached_path.is_file() and cached_path.name.casefold() == filename.casefold():
                return cached_path

        main_path = None
        try:
            main_path = Path(DataClient()._resolver_caminho_excel())
        except DataClientError:
            # A página Ativação pode funcionar com seu próprio arquivo mesmo
            # quando a planilha principal não está disponível.
            pass

        checked: list[Path] = []
        for directory in self._candidate_directories(main_path):
            candidate = directory / filename
            checked.append(candidate)
            if candidate.is_file() and not candidate.name.startswith("~$"):
                self.__class__._caminho_cache = str(candidate)
                return candidate

        found, onedrive_roots = _localizar_excel_no_onedrive(filename)
        if found:
            self.__class__._caminho_cache = str(found)
            return found

        checked_text = "\n".join(f"- {path}" for path in checked)
        onedrive_text = "\n".join(f"- {path}" for path in onedrive_roots)
        raise DataClientError(
            f"O arquivo '{filename}' não foi encontrado.\n\n"
            "Coloque-o ao lado do REDEB2B.exe, na Área de Trabalho, em Documentos "
            "ou em uma pasta sincronizada do OneDrive.\n\n"
            f"Locais diretos verificados:\n{checked_text or '- nenhum'}\n\n"
            f"Pastas do OneDrive verificadas:\n{onedrive_text or '- nenhuma'}",
            status_code=500,
        )

    def _ensure_workbook(self) -> Path:
        path = self._resolve_path()
        if path.is_file():
            return path
        try:
            path.parent.mkdir(parents=True, exist_ok=True)
            self._create_workbook(path)
        except PermissionError as exc:
            raise DataClientError(
                "Não foi possível criar ATIVACAO.xlsx. Verifique a sincronização "
                "e a permissão da pasta do OneDrive.",
                status_code=500,
            ) from exc
        return path

    @staticmethod
    def _create_workbook(path: Path):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "ATIVACAO"
        sheet.append(EXCEL_HEADERS)
        table = Table(displayName="ATIVACAO", ref="A1:M1")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=False,
        )
        sheet.add_table(table)
        sheet.freeze_panes = "A2"
        widths = [13, 28, 22, 27, 15, 17, 12, 18, 18, 12, 27, 14, 14]
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4472C4")
        cities = workbook.create_sheet("CIDADES")
        cities["A1"] = "CIDADES"
        cities["A1"].font = Font(bold=True, color="FFFFFF")
        cities["A1"].fill = PatternFill("solid", fgColor="4472C4")
        cities.column_dimensions["A"].width = 34
        workbook.save(path)
        workbook.close()

    @staticmethod
    def _table(sheet):
        if "ATIVACAO" not in sheet.tables:
            raise DataClientError(
                "A tabela ATIVACAO não foi encontrada na aba ATIVACAO.",
                status_code=500,
            )
        return sheet.tables["ATIVACAO"]

    @classmethod
    def _column_map(cls, sheet, table) -> dict[str, int]:
        min_col, min_row, max_col, _max_row = range_boundaries(table.ref)
        normalized = {
            _normalize_header(sheet.cell(min_row, col).value): col
            for col in range(min_col, max_col + 1)
        }
        mapping = {}
        missing = []
        for field, aliases in HEADER_ALIASES.items():
            match = next((normalized[alias] for alias in aliases if alias in normalized), None)
            if match is None:
                missing.append(field)
            else:
                mapping[field] = match
        if missing:
            readable = ", ".join(REQUIRED_LABELS.get(field, field) for field in missing)
            raise DataClientError(
                f"A tabela ATIVACAO não possui todas as colunas necessárias: {readable}.",
                status_code=500,
            )
        return mapping

    def _load(self):
        path = self._ensure_workbook()
        try:
            workbook = load_workbook(path)
        except PermissionError as exc:
            raise DataClientError(
                "ATIVACAO.xlsx está bloqueado. Feche o arquivo no Excel e tente novamente.",
                status_code=423,
            ) from exc
        if "ATIVACAO" not in workbook.sheetnames:
            workbook.close()
            raise DataClientError("A aba ATIVACAO não foi encontrada no arquivo.", status_code=500)
        sheet = workbook["ATIVACAO"]
        table = self._table(sheet)
        mapping = self._column_map(sheet, table)
        return path, workbook, sheet, table, mapping

    @staticmethod
    def _serialize(field: str, value):
        if field in {"data_agendamento", "data_execucao"}:
            parsed = _parse_date(value)
            return parsed.isoformat() if parsed else ""
        if field == "id":
            try:
                return int(value)
            except (TypeError, ValueError):
                return _text(value)
        return _text(value)

    def _read_rows(self, sheet, table, mapping) -> list[dict]:
        _min_col, min_row, _max_col, max_row = range_boundaries(table.ref)
        rows = []
        for row_number in range(min_row + 1, max_row + 1):
            values = {field: sheet.cell(row_number, column).value for field, column in mapping.items()}
            if not any(value not in (None, "") for value in values.values()):
                continue
            item = {field: self._serialize(field, values[field]) for field in FIELDS}
            item["_row"] = row_number
            rows.append(item)
        return rows

    @staticmethod
    def _save_atomic(path: Path, workbook):
        temporary = None
        try:
            with tempfile.NamedTemporaryFile(
                prefix=".ativacao_", suffix=".xlsx", dir=path.parent, delete=False
            ) as handle:
                temporary = Path(handle.name)
            workbook.save(temporary)
            workbook.close()
            os.replace(temporary, path)
        except PermissionError as exc:
            workbook.close()
            if temporary:
                temporary.unlink(missing_ok=True)
            raise DataClientError(
                "Não foi possível salvar. Feche ATIVACAO.xlsx no Excel e tente novamente.",
                status_code=423,
            ) from exc
        except Exception:
            workbook.close()
            if temporary:
                temporary.unlink(missing_ok=True)
            raise

    @staticmethod
    def _canonical(value, choices, label):
        raw = _text(value)
        if not raw:
            raise DataClientError(f"Campo {label} obrigatório.", status_code=400)
        choice_map = {choice.casefold(): choice for choice in choices}
        if raw.casefold() not in choice_map:
            raise DataClientError(f"Valor inválido para {label}.", status_code=400)
        return choice_map[raw.casefold()]

    def _validate(self, payload: dict, current: dict | None = None) -> dict:
        merged = {field: value for field, value in (current or {}).items() if field in FIELDS}
        merged.update(payload or {})
        result = {}
        result["id"] = _parse_id(
            merged.get("id"),
            required=current is not None or "id" in (payload or {}),
        )
        for field, label in REQUIRED_LABELS.items():
            if field == "data_agendamento":
                result[field] = _parse_date(merged.get(field), required=True, label=label)
            elif not _text(merged.get(field)):
                raise DataClientError(f"Campo {label} obrigatório.", status_code=400)
        result["data_execucao"] = _parse_date(merged.get("data_execucao"), label="Data de Execução")
        for field in ("cliente", "cidade", "servico", "tecnico"):
            result[field] = _text(merged.get(field)).upper()
        result["tecnologia"] = self._canonical(merged.get("tecnologia"), FIXED_OPTIONS["tecnologias"], "Tecnologia")
        result["empresa"] = self._canonical(merged.get("empresa"), FIXED_OPTIONS["empresas"], "Empresa")
        result["status"] = self._canonical(merged.get("status"), FIXED_OPTIONS["status"], "Status")
        result["no_mes"] = self._canonical(merged.get("no_mes"), FIXED_OPTIONS["sim_nao"], "No Mês")
        result["faturado"] = self._canonical(merged.get("faturado"), FIXED_OPTIONS["sim_nao"], "Faturado")
        result["com_rfs"] = self._canonical(merged.get("com_rfs"), FIXED_OPTIONS["sim_nao"], "Com RFS")
        return result

    @staticmethod
    def _excel_value(field, value):
        if field in {"data_agendamento", "data_execucao"}:
            return datetime.combine(value, time.min) if value else None
        return value

    @staticmethod
    def _clean_item(item):
        return {field: item.get(field, "") for field in FIELDS}

    def status_file(self):
        path = self._resolve_path()
        return {"path": str(path), "exists": path.is_file()}

    def _all(self) -> list[dict]:
        with ATIVACAO_LOCK:
            _path, workbook, sheet, table, mapping = self._load()
            try:
                return self._read_rows(sheet, table, mapping)
            finally:
                workbook.close()

    def options(self):
        with ATIVACAO_LOCK:
            _path, workbook, sheet, table, mapping = self._load()
            try:
                rows = self._read_rows(sheet, table, mapping)
                city_values = [item["cidade"] for item in rows]
                if "CIDADES" in workbook.sheetnames:
                    city_sheet = workbook["CIDADES"]
                    for row in city_sheet.iter_rows(values_only=True):
                        for value in row:
                            if value and _normalize_header(value) != "CIDADES":
                                city_values.append(_text(value).upper())
            finally:
                workbook.close()
        return {
            "clientes": _unique_sorted(item["cliente"] for item in rows),
            "cidades": _unique_sorted(city_values),
            "servicos": _unique_sorted(item["servico"] for item in rows),
            **{key: list(values) for key, values in FIXED_OPTIONS.items()},
        }

    def get(self, item_id) -> dict:
        for item in self._all():
            if str(item["id"]) == str(item_id):
                return self._clean_item(item)
        raise DataClientError(f"Registro ID {item_id} não encontrado.", status_code=404)

    def create(self, payload: dict) -> dict:
        normalized = self._validate(payload)
        with ATIVACAO_LOCK:
            path, workbook, sheet, table, mapping = self._load()
            rows = self._read_rows(sheet, table, mapping)
            numeric_ids = [item["id"] for item in rows if isinstance(item["id"], int)]
            if normalized["id"] is None:
                normalized["id"] = (max(numeric_ids) if numeric_ids else 0) + 1
            elif any(str(item["id"]) == str(normalized["id"]) for item in rows):
                workbook.close()
                raise DataClientError(
                    f"Já existe uma atividade com o ID {normalized['id']}.",
                    status_code=409,
                )
            _min_col, min_row, _max_col, max_row = range_boundaries(table.ref)
            used_rows = {item["_row"] for item in rows}
            target_row = next((row for row in range(min_row + 1, max_row + 1) if row not in used_rows), max_row + 1)
            for field in FIELDS:
                cell = sheet.cell(target_row, mapping[field])
                cell.value = self._excel_value(field, normalized[field])
                if field in {"data_agendamento", "data_execucao"}:
                    cell.number_format = "dd/mm/yyyy"
            if target_row > max_row:
                min_col, _min_row, max_col, _max_row = range_boundaries(table.ref)
                table.ref = f"{sheet.cell(min_row, min_col).coordinate}:{sheet.cell(target_row, max_col).coordinate}"
                sheet.auto_filter.ref = table.ref
            self._save_atomic(path, workbook)
        return self.get(normalized["id"])

    def update(self, item_id, payload: dict) -> dict:
        with ATIVACAO_LOCK:
            path, workbook, sheet, table, mapping = self._load()
            rows = self._read_rows(sheet, table, mapping)
            current = next((item for item in rows if str(item["id"]) == str(item_id)), None)
            if current is None:
                workbook.close()
                raise DataClientError(f"Registro ID {item_id} não encontrado.", status_code=404)
            normalized = self._validate(payload, current=current)
            if any(
                item["_row"] != current["_row"]
                and str(item["id"]) == str(normalized["id"])
                for item in rows
            ):
                workbook.close()
                raise DataClientError(
                    f"Já existe uma atividade com o ID {normalized['id']}.",
                    status_code=409,
                )
            for field in FIELDS:
                cell = sheet.cell(current["_row"], mapping[field])
                cell.value = self._excel_value(field, normalized[field])
                if field in {"data_agendamento", "data_execucao"}:
                    cell.number_format = "dd/mm/yyyy"
            self._save_atomic(path, workbook)
        return self.get(normalized["id"])

    def delete(self, item_id) -> dict:
        with ATIVACAO_LOCK:
            path, workbook, sheet, table, mapping = self._load()
            rows = self._read_rows(sheet, table, mapping)
            current = next((item for item in rows if str(item["id"]) == str(item_id)), None)
            if current is None:
                workbook.close()
                raise DataClientError(f"Registro ID {item_id} não encontrado.", status_code=404)
            sheet.delete_rows(current["_row"], 1)
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)
            new_max_row = max(min_row, max_row - 1)
            table.ref = f"{sheet.cell(min_row, min_col).coordinate}:{sheet.cell(new_max_row, max_col).coordinate}"
            sheet.auto_filter.ref = table.ref
            self._save_atomic(path, workbook)
        return {"id": current["id"], "deleted": True}

    @staticmethod
    def _filter(rows: list[dict], filters=None) -> list[dict]:
        filters = filters or {}
        execution_date = _parse_date(filters.get("data_execucao"), label="Data de Execução")
        schedule_date = _parse_date(filters.get("data_agendamento"), label="Data de Agendamento")
        exact_fields = ("servico", "tecnologia", "empresa", "status", "faturado", "com_rfs")
        query = _text(filters.get("q")).casefold()
        result = []
        for item in rows:
            if execution_date and _parse_date(item.get("data_execucao")) != execution_date:
                continue
            if schedule_date and _parse_date(item.get("data_agendamento")) != schedule_date:
                continue
            if any(
                filters.get(field)
                and _text(item.get(field)).casefold() != _text(filters[field]).casefold()
                for field in exact_fields
            ):
                continue
            if query and query not in " ".join(_text(item.get(field)) for field in FIELDS).casefold():
                continue
            result.append(item)
        return result

    def list(self, filters=None, page=1, page_size=20, sort="data_execucao:desc") -> dict:
        rows = self._filter(self._all(), filters)
        field, _, direction = (sort or "data_execucao:desc").partition(":")
        field = field if field in FIELDS else "data_execucao"
        rows.sort(
            key=lambda item: _text(item.get(field)).casefold(),
            reverse=direction.casefold() == "desc",
        )
        rows.sort(key=lambda item: item.get(field) in (None, ""))
        page = max(1, int(page))
        page_size = min(200, max(1, int(page_size)))
        total = len(rows)
        start = (page - 1) * page_size
        return {
            "items": [self._clean_item(item) for item in rows[start:start + page_size]],
            "total": total,
            "page": page,
            "page_size": page_size,
            "total_pages": max(1, (total + page_size - 1) // page_size),
        }

    def export_rows(self, filters=None) -> list[dict]:
        return [self._clean_item(item) for item in self._filter(self._all(), filters)]

    @staticmethod
    def _distribution(rows, field, labels=None, limit=12):
        counts = Counter(_text(item.get(field)) for item in rows if _text(item.get(field)))
        selected = list(labels) if labels is not None else [label for label, _count in counts.most_common(limit)]
        return {"labels": selected, "values": [counts.get(label, 0) for label in selected]}

    def dashboard(self, filters=None) -> dict:
        rows = self._filter(self._all(), filters)
        execution_counts = Counter(
            item["data_execucao"] for item in rows if item.get("data_execucao")
        )
        execution_labels = sorted(execution_counts)
        return {
            "kpis": {
                "total": len(rows),
                "ok": sum(item["status"] == "OK" for item in rows),
                "nok": sum(item["status"] == "NOK" for item in rows),
                "faturado": sum(item["faturado"] == "SIM" for item in rows),
                "com_rfs": sum(item["com_rfs"] == "SIM" for item in rows),
            },
            "por_status": self._distribution(rows, "status", FIXED_OPTIONS["status"]),
            "por_tecnologia": self._distribution(rows, "tecnologia", FIXED_OPTIONS["tecnologias"]),
            "por_servico": self._distribution(rows, "servico"),
            "por_faturado": self._distribution(rows, "faturado", FIXED_OPTIONS["sim_nao"]),
            "por_rfs": self._distribution(rows, "com_rfs", FIXED_OPTIONS["sim_nao"]),
            "por_data_execucao": {
                "labels": execution_labels,
                "values": [execution_counts[label] for label in execution_labels],
            },
        }
