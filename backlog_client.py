# -*- coding: utf-8 -*-
"""Leitura e análises da base compartilhada ``Backlog.xlsx``.

O módulo é deliberadamente somente leitura: a planilha é alimentada por outro
processo e o aplicativo apenas apresenta filtros, indicadores e exportação.
"""

from __future__ import annotations

import math
import os
import re
import sys
import unicodedata
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from excel_client import DataClient, DataClientError
from sharepoint_client import DataStoreError, discover_excel_file


def _normalize(value) -> str:
    text = "" if value is None else str(value).strip().casefold()
    return "".join(
        char for char in unicodedata.normalize("NFKD", text)
        if not unicodedata.combining(char)
    )


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        if value.is_integer():
            return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.isoformat(timespec="minutes") if isinstance(value, datetime) else value.isoformat()
    return str(value).strip()


def _date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = _text(value)
    if not raw:
        return None
    raw = raw.replace("T", " ").split(" ", 1)[0]
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw, pattern).date()
        except ValueError:
            continue
    return None


def _number(value) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return None if isinstance(value, float) and math.isnan(value) else float(value)
    raw = _text(value).replace(" ", "")
    if "," in raw:
        # Formato brasileiro: 1.234,5.
        raw = raw.replace(".", "").replace(",", ".")
    try:
        return float(raw)
    except (TypeError, ValueError):
        return None


class BacklogClient:
    """Localiza e consulta ``Backlog.xlsx`` sem modificar o arquivo."""

    _path_cache: str | None = None

    def __init__(self, excel_path: str | os.PathLike | None = None):
        self._explicit_path = Path(excel_path).expanduser() if excel_path else None

    @staticmethod
    def _candidate_directories(main_path: Path | None = None) -> list[Path]:
        candidates: list[Path] = []
        # O arquivo na mesma biblioteca da base principal tem prioridade.
        if main_path:
            candidates.append(main_path.parent)
        if getattr(sys, "frozen", False):
            candidates.append(Path(sys.executable).resolve().parent)
        else:
            candidates.append(Path(__file__).resolve().parent)
        candidates.append(Path.cwd())
        home = Path.home()
        candidates.extend((home / "Desktop", home / "Área de Trabalho", home / "Documents", home / "Documentos"))
        result, seen = [], set()
        for directory in candidates:
            try:
                key = str(directory.resolve()).casefold()
            except OSError:
                key = str(directory).casefold()
            if key not in seen:
                seen.add(key)
                result.append(directory)
        return result

    def _resolve_path(self) -> Path:
        if self._explicit_path:
            if self._explicit_path.is_file():
                return self._explicit_path
            raise DataClientError(f"O arquivo não foi encontrado: {self._explicit_path}", status_code=404)

        configured = os.getenv("BACKLOG_PATH", "").strip()
        if configured:
            path = Path(configured).expanduser()
            if path.is_file():
                return path
            raise DataClientError(f"O caminho configurado em BACKLOG_PATH não foi encontrado: {path}", status_code=404)

        filename = os.getenv("BACKLOG_FILENAME", "Backlog.xlsx").strip() or "Backlog.xlsx"
        if Path(filename).name != filename or Path(filename).suffix.casefold() != ".xlsx":
            raise DataClientError("BACKLOG_FILENAME deve conter somente o nome de um arquivo .xlsx.", status_code=500)

        cached = self.__class__._path_cache
        if cached:
            path = Path(cached)
            if path.is_file() and path.name.casefold() == filename.casefold():
                return path

        main_path = None
        try:
            main_path = Path(DataClient()._resolver_caminho_excel())
        except DataClientError:
            pass

        checked = []
        for directory in self._candidate_directories(main_path):
            candidate = directory / filename
            checked.append(candidate)
            if candidate.is_file() and not candidate.name.startswith("~$"):
                self.__class__._path_cache = str(candidate)
                return candidate

        try:
            found = discover_excel_file(filename, base_dir=Path(__file__).resolve().parent)
        except DataStoreError as exc:
            raise DataClientError(
                f"O arquivo '{filename}' não foi encontrado ao lado do REDE_B2B.xlsx nem nas pastas sincronizadas do OneDrive.",
                status_code=404,
            ) from exc
        self.__class__._path_cache = str(found)
        return found

    @staticmethod
    def _find_sheet(workbook):
        required = {_normalize("Pedido"), _normalize("Carteira")}
        for sheet in workbook.worksheets:
            headers = {_normalize(cell.value) for cell in sheet[1] if cell.value}
            if required.issubset(headers):
                return sheet
        raise DataClientError("O Backlog.xlsx não possui uma aba com as colunas Pedido e Carteira.", status_code=422)

    def _read(self) -> tuple[list[str], list[dict]]:
        path = self._resolve_path()
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet = self._find_sheet(workbook)
            headers = [_text(cell.value) for cell in sheet[1]]
            rows = []
            for number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                if not any(value not in (None, "") for value in values):
                    continue
                item = {headers[index]: values[index] if index < len(values) else None for index in range(len(headers)) if headers[index]}
                item["_row"] = number
                rows.append(item)
            workbook.close()
            return headers, rows
        except DataClientError:
            raise
        except PermissionError as exc:
            raise DataClientError("O Backlog.xlsx está bloqueado. Feche o arquivo no Excel e tente novamente.", status_code=409) from exc
        except Exception as exc:  # noqa: BLE001
            raise DataClientError(f"Não foi possível ler o Backlog.xlsx: {exc}", status_code=500) from exc

    @staticmethod
    def _value(row: dict, *aliases):
        normalized = {_normalize(key): value for key, value in row.items()}
        for alias in aliases:
            if _normalize(alias) in normalized:
                return normalized[_normalize(alias)]
        return None

    @classmethod
    def _technology(cls, row):
        return cls._value(row, "Tecnologia_Report", "Tecnologia_Venda", "TECNOLOGIA_LinkAtual")

    @classmethod
    def _filter(cls, rows: list[dict], filters: dict | None = None) -> list[dict]:
        filters = filters or {}
        start = _date(filters.get("data_inicio"))
        end = _date(filters.get("data_fim"))
        query = _normalize(filters.get("q"))
        result = []
        for row in rows:
            if filters.get("carteira") and _normalize(cls._value(row, "Carteira")) != _normalize(filters["carteira"]):
                continue
            if filters.get("uf") and _normalize(cls._value(row, "UF")) != _normalize(filters["uf"]):
                continue
            if filters.get("cidade") and _normalize(cls._value(row, "Cidade")) != _normalize(filters["cidade"]):
                continue
            if filters.get("tecnologia") and _normalize(cls._technology(row)) != _normalize(filters["tecnologia"]):
                continue
            entry_date = _date(cls._value(row, "Data_Entrada"))
            if start and (not entry_date or entry_date < start):
                continue
            if end and (not entry_date or entry_date > end):
                continue
            if query:
                searchable = " ".join(_text(cls._value(row, field)) for field in ("Pedido", "ID_Cliente", "ID_Vantive", "Cliente", "Cidade"))
                if query not in _normalize(searchable):
                    continue
            result.append(row)
        return result

    @classmethod
    def _public_row(cls, row: dict) -> dict:
        return {
            "row": row.get("_row"),
            "pedido": _text(cls._value(row, "Pedido")),
            "id_cliente": _text(cls._value(row, "ID_Cliente")),
            "cliente": _text(cls._value(row, "Cliente")),
            "uf": _text(cls._value(row, "UF")),
            "cidade": _text(cls._value(row, "Cidade")),
            "carteira": _text(cls._value(row, "Carteira")),
            "servico": _text(cls._value(row, "Servico")),
            "tecnologia": _text(cls._technology(row)),
            "dias_carteira": _text(cls._value(row, "Dias_CarteiraAtual", "Aging Atual")),
            "data_entrada": _text(cls._value(row, "Data_Entrada"))[:10],
            "data_rfs": _text(cls._value(row, "Data_RFS"))[:10],
        }

    def status(self) -> dict:
        path = self._resolve_path()
        return {"caminho": str(path), "arquivo": path.name, "modo": "Excel compartilhado (somente leitura)"}

    def options(self) -> dict:
        _, rows = self._read()
        unique = lambda values: sorted({_text(value) for value in values if _text(value)}, key=str.casefold)
        return {
            "carteiras": unique(self._value(row, "Carteira") for row in rows),
            "ufs": unique(self._value(row, "UF") for row in rows),
            "cidades": unique(self._value(row, "Cidade") for row in rows),
            "tecnologias": unique(self._technology(row) for row in rows),
        }

    def list_rows(self, filters=None, page=1, page_size=20) -> dict:
        _, rows = self._read()
        rows = self._filter(rows, filters)
        rows.sort(key=lambda row: (_number(self._value(row, "Dias_CarteiraAtual", "Aging Atual")) or -1), reverse=True)
        page = max(1, int(page or 1))
        page_size = min(100, max(10, int(page_size or 20)))
        total = len(rows)
        start = (page - 1) * page_size
        return {"items": [self._public_row(row) for row in rows[start:start + page_size]], "total": total, "page": page, "page_size": page_size}

    def dashboard(self, filters=None) -> dict:
        _, rows = self._read()
        rows = self._filter(rows, filters)
        total = len(rows)
        rfs_rows = [row for row in rows if _date(self._value(row, "Data_RFS"))]
        aging = [_number(self._value(row, "Dias_CarteiraAtual", "Aging Atual")) for row in rows]
        aging = [value for value in aging if value is not None]

        carteira_counts = Counter(_text(self._value(row, "Carteira")) or "Não informado" for row in rows)
        carteira_aging: dict[str, list[float]] = {}
        for row in rows:
            name = _text(self._value(row, "Carteira")) or "Não informado"
            value = _number(self._value(row, "Dias_CarteiraAtual", "Aging Atual"))
            if value is not None:
                carteira_aging.setdefault(name, []).append(value)
        carteiras = []
        for name, count in sorted(carteira_counts.items(), key=lambda item: (-item[1], item[0].casefold())):
            values = carteira_aging.get(name, [])
            average = round(sum(values) / len(values), 1) if values else None
            level = "ok" if average is not None and average <= 15 else "warning" if average is not None and average <= 30 else "critical" if average is not None else "neutral"
            carteiras.append({"nome": name, "quantidade": count, "aging_medio": average, "nivel": level})

        rfs_by_day = Counter(_date(self._value(row, "Data_RFS")) for row in rfs_rows)
        dates = sorted(value for value in rfs_by_day if value)
        timeline_dates = dates[-30:]
        latest_updates = [_date(self._value(row, "Data_Atualizacao")) for row in rows]
        latest_updates = [value for value in latest_updates if value]
        return {
            "total": total,
            "com_rfs": len(rfs_rows),
            "percentual_rfs": round((len(rfs_rows) / total * 100), 1) if total else 0,
            "aging_medio": round(sum(aging) / len(aging), 1) if aging else 0,
            "ultima_atualizacao": max(latest_updates).isoformat() if latest_updates else "",
            "carteiras": carteiras,
            "rfs_por_dia": {
                "labels": [value.isoformat() for value in timeline_dates],
                "values": [rfs_by_day[value] for value in timeline_dates],
            },
        }

    def export_rows(self, filters=None) -> tuple[list[str], list[list]]:
        headers, rows = self._read()
        rows = self._filter(rows, filters)
        return headers, [[row.get(header) for header in headers] for row in rows]
