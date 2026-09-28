"""Acesso compartilhado ao arquivo REPARO.xlsx.

O arquivo é localizado pelos mesmos critérios das demais bases do aplicativo:
ao lado do executável, Área de Trabalho, Documentos ou OneDrive.
"""

from __future__ import annotations

import os
import sys
import tempfile
import threading
import unicodedata
from collections import Counter, defaultdict
from copy import copy
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.cell import range_boundaries

from excel_client import DataClient, DataClientError, _localizar_excel_no_onedrive


REPARO_LOCK = threading.RLock()


class _SheetRange:
    """Faixa de dados usada internamente sem criar uma Tabela do Excel.

    A REPARO.xlsx pode ser editada no Excel/OneDrive e não depende de recursos
    estruturados. Manter somente o AutoFiltro da planilha evita que o Excel
    tente reparar repetidamente /xl/tables/table1.xml.
    """

    def __init__(self, ref: str):
        self.ref = ref

FIELDS = [
    "BD", "ID_VANTIVE", "STATUS", "RECLAMACAO", "CLIENTE", "ENDERECO",
    "CIDADE", "UF", "CLUSTER", "LP_15", "DATA_ABERTURA",
    "DATA_ENCERRAMENTO", "BAIXA_CODIGO", "GRUPO_BAIXA", "USUARIO_BAIXA",
    "REINC_30D", "REINC_TIPO", "KPI", "TMR", "TEMPO_PARADA", "EPS",
    "TIPO_DEFEITO", "PARADA_RELOGIO", "INICIO", "FIM",
    "RECLAMACAO_CLIENTE", "DESCRICAO_FALHA", "MOTIVO_REAL_BD", "ARD_ERB",
    "TECNICO", "FILA_ENCERRAMENTO", "TECNOLOGIA", "TMR_PARADA",
    "TIPO_ARD_ERB", "ATUALIZADO_EM",
]

EXCEL_HEADERS = list(FIELDS)
DATE_FIELDS = {"DATA_ABERTURA", "DATA_ENCERRAMENTO"}
DATETIME_FIELDS = {"INICIO", "FIM", "ATUALIZADO_EM"}
REQUIRED = {"ID_VANTIVE": "ID"}

EPS_OPTIONS = {"VIVO", "RS TELECOM", "STAFF"}
TIPO_DEFEITO_OPTIONS = {"CAMPO", "SISTEMICO"}
SIM_NAO_OPTIONS = {"SIM", "NÃO"}
TIPO_ARD_ERB_OPTIONS = {"ERB", "ARD"}

LIST_SHEET_NAME = "LISTAS"
LIST_HEADERS = {
    "motivo_real_bd": "MOTIVO REAL DO BD",
    "tecnico": "TÉCNICO",
    "fila_encerramento": "FILA DE ENCERRAMENTO",
    "tecnologia": "TECNOLOGIA",
    "usuario": "USUÁRIO",
}

HEADER_ALIASES = {
    "ID_VANTIVE": {"ID"},
    "TIPO_DEFEITO": {"TIPO DE DEFEITO"},
    "PARADA_RELOGIO": {"PARADA DE RELOGIO"},
    "RECLAMACAO_CLIENTE": {"RECLAMACAO DO CLIENTE"},
    "DESCRICAO_FALHA": {"DESCRICAO DA FALHA"},
    "MOTIVO_REAL_BD": {"MOTIVO REAL DO BD"},
    "FILA_ENCERRAMENTO": {"FILA DE ENCERRAMENTO"},
    "TMR_PARADA": {"TMR (PARADA)", "TMR-(PARADA)"},
    "TIPO_ARD_ERB": {"TIPO ARD/ERB", "TIPO ARD ERB"},
    "ARD_ERB": {"ARD/ERB"},
    "USUARIO_BAIXA": {"USUARIO"},
    "ATUALIZADO_EM": {"DATA ATUALIZACAO", "ATUALIZADO EM"},
    "RECLAMACAO": {"RECLAMACAO"},
}


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _normalize(value) -> str:
    text = _text(value).upper().replace("_", " ")
    text = "".join(
        char for char in unicodedata.normalize("NFKD", text)
        if not unicodedata.combining(char)
    )
    return " ".join(text.replace("-", " ").split())


def _parse_date(value, *, required=False, label="Data") -> date | None:
    if value is None or _text(value) == "":
        if required:
            raise DataClientError(f"Campo {label} obrigatório.", status_code=400)
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = _text(value)
    for pattern in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(raw[:10], pattern).date()
        except ValueError:
            continue
    raise DataClientError(f"{label} inválida.", status_code=400)


def _parse_datetime(value, *, required=False, label="Data e hora") -> datetime | None:
    if value is None or _text(value) == "":
        if required:
            raise DataClientError(f"Campo {label} obrigatório.", status_code=400)
        return None
    if isinstance(value, datetime):
        return value.replace(second=0, microsecond=0)
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    raw = _text(value)
    normalized = raw.replace("Z", "").replace("T", " ")
    for pattern in (
        "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M",
        "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M",
    ):
        try:
            return datetime.strptime(normalized, pattern).replace(second=0, microsecond=0)
        except ValueError:
            continue
    raise DataClientError(f"{label} inválida.", status_code=400)


def _parse_date_or_datetime(value, *, required=False, label="Data") -> date | datetime | None:
    if isinstance(value, datetime):
        return value.replace(second=0, microsecond=0)
    raw = _text(value)
    if raw and ("T" in raw or ":" in raw):
        return _parse_datetime(value, required=required, label=label)
    return _parse_date(value, required=required, label=label)


def _as_datetime(value) -> datetime | None:
    parsed = _parse_date_or_datetime(value)
    if parsed is None:
        return None
    return parsed if isinstance(parsed, datetime) else datetime.combine(parsed, time.min)


def _calculate_tmr_parada(data_abertura, data_encerramento, inicio, fim) -> str:
    opened = _as_datetime(data_abertura)
    closed = _as_datetime(data_encerramento)
    if not opened or not closed:
        return ""
    total_seconds = (closed - opened).total_seconds()
    if total_seconds < 0:
        raise DataClientError(
            "Data de encerramento não pode ser anterior à data de abertura.",
            status_code=400,
        )
    paused_seconds = 0
    pause_start = _parse_datetime(inicio, label="Início")
    pause_end = _parse_datetime(fim, label="Fim")
    if bool(pause_start) != bool(pause_end):
        raise DataClientError(
            "Preencha Início e Fim para calcular a parada do relógio.",
            status_code=400,
        )
    if pause_start and pause_end:
        paused_seconds = (pause_end - pause_start).total_seconds()
        if paused_seconds < 0:
            raise DataClientError("Fim não pode ser anterior ao Início.", status_code=400)
    hours = max(0, total_seconds - paused_seconds) / 3600
    return f"{hours:.2f}"


def _number(value) -> float | None:
    raw = _text(value).replace(",", ".")
    try:
        return float(raw)
    except (TypeError, ValueError):
        return None


def _unique(values):
    result = {_text(value) for value in values if _text(value)}
    return sorted(result, key=str.casefold)


class ReparoClient:
    """CRUD, filtros e indicadores da base REPARO."""

    _caminho_cache: str | None = None

    def __init__(self, excel_path: str | os.PathLike | None = None):
        self._explicit_path = Path(excel_path).expanduser() if excel_path else None

    @staticmethod
    def _candidate_directories(main_path: Path | None = None) -> list[Path]:
        candidates = []
        if getattr(sys, "frozen", False):
            candidates.append(Path(sys.executable).resolve().parent)
        else:
            candidates.append(Path(__file__).resolve().parent)
        candidates.append(Path.cwd())
        home = Path.home()
        candidates.extend((
            home / "Desktop", home / "Área de Trabalho",
            home / "Documents", home / "Documentos",
        ))
        if main_path:
            candidates.append(main_path.parent)
        result, seen = [], set()
        for directory in candidates:
            try:
                key = str(directory.resolve()).casefold()
            except OSError:
                key = str(directory).casefold()
            if key not in seen:
                seen.add(key)
                result.append(directory)
        return result

    def _resolve_path(self) -> Path:
        if self._explicit_path:
            return self._explicit_path
        configured = os.getenv("REPARO_PATH", "").strip()
        if configured:
            return Path(configured).expanduser()
        filename = os.getenv("REPARO_FILENAME", "REPARO.xlsx").strip() or "REPARO.xlsx"
        if Path(filename).name != filename or Path(filename).suffix.casefold() != ".xlsx":
            raise DataClientError(
                "REPARO_FILENAME deve conter somente o nome de um arquivo .xlsx.",
                status_code=500,
            )
        cached = self.__class__._caminho_cache
        if cached:
            path = Path(cached)
            if path.is_file() and path.name.casefold() == filename.casefold():
                return path

        main_path = None
        try:
            main_path = Path(DataClient()._resolver_caminho_excel())
        except DataClientError:
            pass

        checked = []
        for directory in self._candidate_directories(main_path):
            candidate = directory / filename
            checked.append(candidate)
            if candidate.is_file() and not candidate.name.startswith("~$"):
                self.__class__._caminho_cache = str(candidate)
                return candidate

        found, onedrive_roots = _localizar_excel_no_onedrive(filename)
        if found:
            self.__class__._caminho_cache = str(found)
            return found
        raise DataClientError(
            f"O arquivo '{filename}' não foi encontrado. Coloque-o ao lado do "
            "REDEB2B.exe, na Área de Trabalho, em Documentos ou em uma pasta "
            "sincronizada do OneDrive.\n\nLocais verificados:\n"
            + "\n".join(f"- {path}" for path in [*checked, *onedrive_roots]),
            status_code=500,
        )

    def _ensure_workbook(self) -> Path:
        path = self._resolve_path()
        if path.is_file():
            return path
        path.parent.mkdir(parents=True, exist_ok=True)
        self._create_workbook(path)
        return path

    @staticmethod
    def _create_workbook(path: Path):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "REPARO"
        sheet.append(EXCEL_HEADERS)
        sheet.freeze_panes = "A2"
        for index, header in enumerate(EXCEL_HEADERS, 1):
            cell = sheet.cell(1, index)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2C7A72")
            sheet.column_dimensions[cell.column_letter].width = max(14, min(28, len(header) + 3))
        workbook.save(path)
        workbook.close()

    @classmethod
    def _find_sheet(cls, workbook):
        if "REPARO" in workbook.sheetnames:
            return workbook["REPARO"]
        required = {_normalize(field) for field in REQUIRED}
        for sheet in workbook.worksheets:
            headers = {_normalize(cell.value) for cell in sheet[1] if cell.value}
            if required.issubset(headers):
                return sheet
        return workbook.active

    @classmethod
    def _prepare_sheet(cls, sheet):
        changed = False
        header_map = {
            _normalize(sheet.cell(1, column).value): column
            for column in range(1, max(sheet.max_column, 1) + 1)
            if sheet.cell(1, column).value
        }
        next_column = sheet.max_column + 1 if header_map else 1
        for field in FIELDS:
            aliases = {_normalize(field), *(_normalize(value) for value in HEADER_ALIASES.get(field, set()))}
            if not any(alias in header_map for alias in aliases):
                column = next_column
                next_column += 1
                source = sheet.cell(1, max(1, column - 1))
                target = sheet.cell(1, column, field)
                if column > 1:
                    target._style = copy(source._style)
                target.font = copy(target.font)
                sheet.column_dimensions[target.column_letter].width = max(14, min(28, len(field) + 3))
                header_map[_normalize(field)] = column
                changed = True

        max_row = max(sheet.max_row, 1)
        max_col = max(sheet.max_column, len(FIELDS))
        ref = f"A1:{sheet.cell(max_row, max_col).coordinate}"
        if sheet.tables:
            for table_name in list(sheet.tables):
                del sheet.tables[table_name]
            changed = True
        filter_ref = ref if max_row > 1 else None
        if sheet.auto_filter.ref != filter_ref:
            sheet.auto_filter.ref = filter_ref
            changed = True
        sheet.freeze_panes = sheet.freeze_panes or "A2"
        return _SheetRange(ref), changed

    @classmethod
    def _column_map(cls, sheet, table):
        min_col, min_row, max_col, _ = range_boundaries(table.ref)
        normalized = {
            _normalize(sheet.cell(min_row, column).value): column
            for column in range(min_col, max_col + 1)
        }
        mapping = {}
        for field in FIELDS:
            aliases = [_normalize(field), *[_normalize(value) for value in HEADER_ALIASES.get(field, set())]]
            mapping[field] = next(normalized[alias] for alias in aliases if alias in normalized)
        return mapping

    @staticmethod
    def _prepare_lists_sheet(workbook):
        changed = False
        if LIST_SHEET_NAME not in workbook.sheetnames:
            sheet = workbook.create_sheet(LIST_SHEET_NAME)
            changed = True
        else:
            sheet = workbook[LIST_SHEET_NAME]

        existing = {
            _normalize(sheet.cell(1, column).value): column
            for column in range(1, max(sheet.max_column, 1) + 1)
            if sheet.cell(1, column).value
        }
        next_column = max(sheet.max_column, 0) + 1 if existing else 1
        for header in LIST_HEADERS.values():
            key = _normalize(header)
            if key in existing:
                continue
            column = next_column
            next_column += 1
            cell = sheet.cell(1, column, header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2C7A72")
            sheet.column_dimensions[cell.column_letter].width = max(20, len(header) + 4)
            existing[key] = column
            changed = True
        sheet.freeze_panes = "A2"
        return changed

    @staticmethod
    def _read_list_options(workbook):
        sheet = workbook[LIST_SHEET_NAME]
        columns = {
            _normalize(sheet.cell(1, column).value): column
            for column in range(1, max(sheet.max_column, 1) + 1)
            if sheet.cell(1, column).value
        }
        options = {}
        for key, header in LIST_HEADERS.items():
            column = columns.get(_normalize(header))
            values = [] if column is None else [
                sheet.cell(row, column).value
                for row in range(2, sheet.max_row + 1)
            ]
            options[key] = _unique(values)
        return options

    def _load(self):
        path = self._ensure_workbook()
        try:
            workbook = load_workbook(path)
        except PermissionError as exc:
            raise DataClientError(
                "REPARO.xlsx está bloqueado. Feche o arquivo no Excel e tente novamente.",
                status_code=423,
            ) from exc
        sheet = self._find_sheet(workbook)
        table, changed = self._prepare_sheet(sheet)
        lists_changed = self._prepare_lists_sheet(workbook)
        if changed or lists_changed:
            self._save_atomic(path, workbook)
            workbook = load_workbook(path)
            sheet = self._find_sheet(workbook)
            table, _ = self._prepare_sheet(sheet)
        return path, workbook, sheet, table, self._column_map(sheet, table)

    @staticmethod
    def _save_atomic(path: Path, workbook):
        temporary = None
        try:
            with tempfile.NamedTemporaryFile(
                prefix=".reparo_", suffix=".xlsx", dir=path.parent, delete=False
            ) as handle:
                temporary = Path(handle.name)
            workbook.save(temporary)
            workbook.close()
            os.replace(temporary, path)
        except PermissionError as exc:
            workbook.close()
            if temporary:
                temporary.unlink(missing_ok=True)
            raise DataClientError(
                "Não foi possível salvar. Feche REPARO.xlsx no Excel e tente novamente.",
                status_code=423,
            ) from exc
        except Exception:
            workbook.close()
            if temporary:
                temporary.unlink(missing_ok=True)
            raise

    @staticmethod
    def _serialize(field, value):
        if field in DATE_FIELDS:
            parsed = _parse_date_or_datetime(value)
            if isinstance(parsed, datetime):
                if parsed.time() == time.min:
                    return parsed.date().isoformat()
                return parsed.isoformat(timespec="minutes")
            return parsed.isoformat() if parsed else ""
        if field in DATETIME_FIELDS:
            parsed = _parse_datetime(value)
            return parsed.isoformat(timespec="minutes") if parsed else ""
        return _text(value)

    def _read_rows(self, sheet, table, mapping):
        _, min_row, _, max_row = range_boundaries(table.ref)
        rows = []
        for row_number in range(min_row + 1, max_row + 1):
            values = {field: sheet.cell(row_number, column).value for field, column in mapping.items()}
            if not any(value not in (None, "") for value in values.values()):
                continue
            item = {field: self._serialize(field, values[field]) for field in FIELDS}
            item["_row"] = row_number
            rows.append(item)
        return rows

    def _all(self):
        with REPARO_LOCK:
            _path, workbook, sheet, table, mapping = self._load()
            try:
                return self._read_rows(sheet, table, mapping)
            finally:
                workbook.close()

    @staticmethod
    def _clean(item):
        result = {field: item.get(field, "") for field in FIELDS}
        result["row_number"] = item.get("_row")
        return result

    @staticmethod
    def _find(rows, item_id, row_number=None):
        target_row = None
        if row_number not in (None, ""):
            try:
                target_row = int(row_number)
            except (TypeError, ValueError) as exc:
                raise DataClientError("Referência da linha inválida.", status_code=400) from exc
        return next((
            item for item in rows
            if item["ID_VANTIVE"].casefold() == _text(item_id).casefold()
            and (target_row is None or item["_row"] == target_row)
        ), None)

    @staticmethod
    def _validate(payload, current=None):
        merged = {field: (current or {}).get(field, "") for field in FIELDS}
        merged.update({key: value for key, value in (payload or {}).items() if key in FIELDS})
        for field, label in REQUIRED.items():
            if not _text(merged.get(field)):
                raise DataClientError(f"Campo {label} obrigatório.", status_code=400)

        normalized_choices = (
            ("EPS", EPS_OPTIONS, "EPS"),
            ("TIPO_DEFEITO", TIPO_DEFEITO_OPTIONS, "Tipo de defeito"),
            ("PARADA_RELOGIO", SIM_NAO_OPTIONS, "Parada de relógio"),
            ("TIPO_ARD_ERB", TIPO_ARD_ERB_OPTIONS, "Tipo ERB/ARD"),
        )
        for field, choices, label in normalized_choices:
            value = _text(merged.get(field)).upper()
            if value and value not in choices:
                raise DataClientError(f"Valor inválido para {label}.", status_code=400)
            merged[field] = value

        if len(_text(merged.get("ARD_ERB"))) > 10:
            raise DataClientError(
                "ARD/ERB deve ter no máximo 10 caracteres.", status_code=400
            )

        result = {field: _text(merged.get(field)) for field in FIELDS}
        for field in DATE_FIELDS:
            parsed = _parse_date_or_datetime(
                merged.get(field), label=field.replace("_", " ").title()
            )
            result[field] = parsed
        for field in {"INICIO", "FIM"}:
            result[field] = _parse_datetime(
                merged.get(field), label=field.replace("_", " ").title()
            )
        result["TMR_PARADA"] = _calculate_tmr_parada(
            result["DATA_ABERTURA"], result["DATA_ENCERRAMENTO"],
            result["INICIO"], result["FIM"],
        )
        result["ATUALIZADO_EM"] = datetime.now().replace(second=0, microsecond=0)
        return result

    @staticmethod
    def _excel_value(field, value):
        if field in DATE_FIELDS | DATETIME_FIELDS:
            if not value:
                return None
            return value if isinstance(value, datetime) else datetime.combine(value, time.min)
        return value

    def status_file(self):
        path = self._resolve_path()
        return {"path": str(path), "exists": path.is_file()}

    def options(self):
        with REPARO_LOCK:
            _path, workbook, sheet, table, mapping = self._load()
            try:
                rows = self._read_rows(sheet, table, mapping)
                configured = self._read_list_options(workbook)
            finally:
                workbook.close()
        return {
            "status": _unique([item["STATUS"] for item in rows] + ["Aberto", "Em Andamento", "Encerrado"]),
            "uf": _unique(item["UF"] for item in rows),
            "eps": ["VIVO", "RS TELECOM", "STAFF"],
            "tipo_defeito": ["CAMPO", "SISTEMICO"],
            "parada_relogio": ["SIM", "NÃO"],
            "tipo_ard_erb": ["ERB", "ARD"],
            **configured,
        }

    def get(self, item_id, row_number=None):
        item = self._find(self._all(), item_id, row_number)
        if not item:
            raise DataClientError(f"Registro ID {item_id} não encontrado.", status_code=404)
        return self._clean(item)

    def create(self, payload):
        normalized = self._validate(payload)
        with REPARO_LOCK:
            path, workbook, sheet, table, mapping = self._load()
            rows = self._read_rows(sheet, table, mapping)
            if any(item["ID_VANTIVE"].casefold() == normalized["ID_VANTIVE"].casefold() for item in rows):
                workbook.close()
                raise DataClientError(
                    f"Já existe um reparo com o ID Vantive {normalized['ID_VANTIVE']}.",
                    status_code=409,
                )
            _, min_row, _, max_row = range_boundaries(table.ref)
            target_row = max(max_row + 1, min_row + 1)
            for field, column in mapping.items():
                cell = sheet.cell(target_row, column, self._excel_value(field, normalized[field]))
                if field in DATE_FIELDS:
                    cell.number_format = "dd/mm/yyyy hh:mm"
                elif field in DATETIME_FIELDS:
                    cell.number_format = "dd/mm/yyyy hh:mm"
            min_col, _, max_col, _ = range_boundaries(table.ref)
            table.ref = f"{sheet.cell(min_row, min_col).coordinate}:{sheet.cell(target_row, max_col).coordinate}"
            sheet.auto_filter.ref = table.ref
            self._save_atomic(path, workbook)
        return self.get(normalized["ID_VANTIVE"], target_row)

    def update(self, item_id, payload, row_number=None):
        with REPARO_LOCK:
            path, workbook, sheet, table, mapping = self._load()
            rows = self._read_rows(sheet, table, mapping)
            current = self._find(rows, item_id, row_number)
            if not current:
                workbook.close()
                raise DataClientError(f"Registro ID {item_id} não encontrado.", status_code=404)
            normalized = self._validate(payload, current)
            if normalized["ID_VANTIVE"].casefold() != current["ID_VANTIVE"].casefold() and any(
                item["_row"] != current["_row"]
                and item["ID_VANTIVE"].casefold() == normalized["ID_VANTIVE"].casefold()
                for item in rows
            ):
                workbook.close()
                raise DataClientError(
                    f"Já existe um reparo com o ID Vantive {normalized['ID_VANTIVE']}.",
                    status_code=409,
                )
            for field, column in mapping.items():
                cell = sheet.cell(current["_row"], column, self._excel_value(field, normalized[field]))
                if field in DATE_FIELDS:
                    cell.number_format = "dd/mm/yyyy hh:mm"
                elif field in DATETIME_FIELDS:
                    cell.number_format = "dd/mm/yyyy hh:mm"
            self._save_atomic(path, workbook)
        return self.get(normalized["ID_VANTIVE"], current["_row"])

    def delete(self, item_id, row_number=None):
        with REPARO_LOCK:
            path, workbook, sheet, table, mapping = self._load()
            rows = self._read_rows(sheet, table, mapping)
            current = self._find(rows, item_id, row_number)
            if not current:
                workbook.close()
                raise DataClientError(f"Registro ID {item_id} não encontrado.", status_code=404)
            sheet.delete_rows(current["_row"], 1)
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)
            new_max_row = max(min_row, max_row - 1)
            table.ref = f"{sheet.cell(min_row, min_col).coordinate}:{sheet.cell(new_max_row, max_col).coordinate}"
            sheet.auto_filter.ref = table.ref if new_max_row > min_row else None
            self._save_atomic(path, workbook)
        return {"id": item_id, "deleted": True}

    @staticmethod
    def _filter(rows, filters):
        filters = filters or {}
        result = rows
        for field, key in (("STATUS", "status"), ("UF", "uf"), ("EPS", "eps"), ("TECNOLOGIA", "tecnologia")):
            wanted = _text(filters.get(key)).casefold()
            if wanted:
                result = [item for item in result if item[field].casefold() == wanted]
        start = _parse_date(filters.get("de"), label="Data inicial") if filters.get("de") else None
        end = _parse_date(filters.get("ate"), label="Data final") if filters.get("ate") else None
        if start and end and start > end:
            raise DataClientError("A data inicial não pode ser maior que a data final.", status_code=400)
        if start:
            result = [item for item in result if item["DATA_ABERTURA"] and item["DATA_ABERTURA"] >= start.isoformat()]
        if end:
            result = [item for item in result if item["DATA_ABERTURA"] and item["DATA_ABERTURA"] <= end.isoformat()]
        query = _text(filters.get("busca")).casefold()
        if query:
            search_fields = ("CLIENTE", "RECLAMACAO_CLIENTE", "DESCRICAO_FALHA", "ID_VANTIVE", "BD")
            result = [
                item for item in result
                if query in " ".join(item[field] for field in search_fields).casefold()
            ]
        return result

    def list(self, filters=None, page=1, page_size=25, sort="DATA_ABERTURA:desc"):
        rows = self._filter(self._all(), filters)
        field, _, direction = (sort or "DATA_ABERTURA:desc").partition(":")
        field = field if field in FIELDS else "DATA_ABERTURA"
        rows.sort(key=lambda item: (not bool(item[field]), item[field].casefold()), reverse=direction.lower() == "desc")
        page = max(1, int(page))
        page_size = max(1, min(5000, int(page_size)))
        start = (page - 1) * page_size
        return {
            "items": [self._clean(item) for item in rows[start:start + page_size]],
            "total": len(rows), "page": page, "page_size": page_size,
        }

    def export_rows(self, filters=None):
        return [self._clean(item) for item in self._filter(self._all(), filters)]

    def dashboard(self, filters=None):
        rows = self._filter(self._all(), filters)
        total = len(rows)
        encerrados = sum("encerr" in item["STATUS"].casefold() for item in rows)
        tmrs = [value for value in (_number(item["TMR"]) for item in rows) if value is not None]
        reincidentes = sum(item["REINC_30D"].casefold().startswith("s") for item in rows)

        by_month = Counter(item["DATA_ABERTURA"][:7] for item in rows if item["DATA_ABERTURA"])
        by_status = Counter(item["STATUS"] or "Sem status" for item in rows)
        by_reason = Counter(item["MOTIVO_REAL_BD"] for item in rows if item["MOTIVO_REAL_BD"])
        by_technology = defaultdict(list)
        for item in rows:
            value = _number(item["TMR"])
            if item["TECNOLOGIA"] and value is not None:
                by_technology[item["TECNOLOGIA"]].append(value)
        reasons = by_reason.most_common(10)
        technologies = sorted(by_technology, key=str.casefold)

        period_end = date.today()
        period_start = period_end - timedelta(days=29)
        repairs_by_lp = {}
        for item in rows:
            lp_15 = _text(item.get("LP_15"))
            if not lp_15:
                continue
            try:
                opened_at = _parse_date(item.get("DATA_ABERTURA"))
            except DataClientError:
                continue
            if not opened_at or opened_at < period_start or opened_at > period_end:
                continue
            key = lp_15.casefold()
            group = repairs_by_lp.setdefault(key, {
                "lp_15": lp_15,
                "cliente": "",
                "quantidade": 0,
                "motivos": Counter(),
                "ultima_data": date.min,
            })
            group["quantidade"] += 1
            reason = _text(item.get("MOTIVO_REAL_BD")) or "Não informado"
            group["motivos"][reason] += 1
            if opened_at >= group["ultima_data"]:
                group["ultima_data"] = opened_at
                group["lp_15"] = lp_15
                if _text(item.get("CLIENTE")):
                    group["cliente"] = _text(item.get("CLIENTE"))

        repeated_lp = []
        for group in repairs_by_lp.values():
            if group["quantidade"] < 2:
                continue
            repeated_lp.append({
                "lp_15": group["lp_15"],
                "cliente": group["cliente"],
                "quantidade": group["quantidade"],
                "motivos": [
                    {"motivo": reason, "quantidade": quantity}
                    for reason, quantity in group["motivos"].most_common()
                ],
            })
        repeated_lp.sort(key=lambda item: (-item["quantidade"], item["lp_15"].casefold()))

        return {
            "kpis": {
                "total": total,
                "abertos": total - encerrados,
                "encerrados": encerrados,
                "tmr_medio": round(sum(tmrs) / len(tmrs), 2) if tmrs else 0,
                "taxa_reincidencia": round(reincidentes / total * 100, 1) if total else 0,
            },
            "serie": {"labels": sorted(by_month), "values": [by_month[key] for key in sorted(by_month)]},
            "status": {"labels": list(by_status), "values": list(by_status.values())},
            "motivos": {"labels": [key for key, _ in reasons], "values": [value for _, value in reasons]},
            "tecnologias": {
                "labels": technologies,
                "values": [round(sum(by_technology[key]) / len(by_technology[key]), 2) for key in technologies],
            },
            "reincidencias_lp_30d": {
                "data_inicio": period_start.isoformat(),
                "data_fim": period_end.isoformat(),
                "itens": repeated_lp,
            },
        }
