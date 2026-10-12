import tempfile
import unittest
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook

from analitico_client import AnaliticoClient, _hours


HEADERS = [
    "BD", "TA", "LP15", "KPI", "ABERTURA", "ENCERRAMENTO", "REGIONAL",
    "UF", "CLUSTER", "CIDADE", "CLIENTE", "GRUPO BAIXA", "USUARIO BAIXA",
    "ARMÁRIO", "ERB", "TECNOLOGIA", "R30 TRATATIVAS", "MASSIVA",
    "TEMPO CORRIDO", "TEMPO PARADA", "CLIENTE CRÍTICO",
]


class AnaliticoTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name) / "ANALITICO.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "ANALITICO"
        sheet.append(HEADERS)
        today = date.today()
        sheet.append(["BD1", "", "LP-1", "", today, today, "SUL", "PR", "C1", "CURITIBA", "CLIENTE A", "", "", "ARM-1", "ERB-1", "GPON", "SIM", "NÃO", 10, 2, "SIM"])
        sheet.append(["BD2", "", "LP-1", "", today - timedelta(days=5), today, "SUL", "PR", "C1", "CURITIBA", "CLIENTE A", "", "", "ARM-1", "ERB-1", "GPON", "X", "NÃO", "08:00", "01:30", "SIM"])
        sheet.append(["BD3", "", "LP-2", "", today - timedelta(days=20), today, "SUL", "PR", "C2", "PINHAIS", "CLIENTE B", "", "", "ARM-2", "ERB-2", "ERB", "", "NÃO", 6, 1, "SIM"])
        workbook.save(self.path)
        workbook.close()
        self.client = AnaliticoClient(self.path)

    def tearDown(self):
        self.temp.cleanup()

    def test_shared_analytics_and_critical_windows(self):
        today = date.today()
        data = self.client.dashboard(start=today - timedelta(days=29), end=today)
        self.assertEqual(data["total_periodo"], 3)
        self.assertEqual(sum(data["bds_por_mes"]["values"]), 3)
        self.assertEqual(data["armarios_criticos"][0], {"nome": "ARM-1", "quantidade": 2})
        self.assertEqual(data["erbs_criticas"][0], {"nome": "ERB-1", "quantidade": 2})
        self.assertEqual(data["r30_tratativas"][0]["quantidade"], 2)
        windows = {item["dias"]: item["itens"] for item in data["clientes_criticos"]}
        self.assertEqual(windows[30][0]["quantidade"], 2)
        self.assertEqual(windows[7][0]["quantidade"], 2)
        self.assertEqual(data["tmr_medio_horas"], 6.5)

    def test_technology_filter_and_duration_conversion(self):
        today = date.today()
        data = self.client.dashboard(start=today - timedelta(days=29), end=today, technologies=["ERB"])
        self.assertEqual(data["total_periodo"], 1)
        self.assertEqual(data["tmr_medio_horas"], 5)
        self.assertEqual(_hours(0.5), 12)
        self.assertEqual(_hours("02:30"), 2.5)

    def test_rankings_are_limited_to_top_ten(self):
        workbook = load_workbook(self.path)
        sheet = workbook["ANALITICO"]
        today = date.today()
        for index in range(12):
            for repeat in range(index + 2):
                sheet.append([
                    f"EX-{index}-{repeat}", "", f"LP-{index + 10}", "", today, today,
                    "SUL", "PR", "C", "CURITIBA", f"CLIENTE {index}", "", "",
                    f"ARM-{index + 10}", f"ERB-{index + 10}", "GPON", "SIM", "NÃO",
                    2, 0, "SIM",
                ])
        workbook.save(self.path)
        workbook.close()

        data = self.client.dashboard(start=today, end=today)
        self.assertEqual(len(data["armarios_criticos"]), 10)
        self.assertEqual(len(data["erbs_criticas"]), 10)
        self.assertEqual(len(data["r30_tratativas"]), 10)
        self.assertTrue(all(len(window["itens"]) <= 10 for window in data["clientes_criticos"]))


if __name__ == "__main__":
    unittest.main()
