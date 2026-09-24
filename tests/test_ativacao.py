import os
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import patch
from urllib.parse import quote

from openpyxl import load_workbook

import api
from app import create_app
from ativacao_client import AtivacaoClient, EXCEL_HEADERS
from excel_client import DataClientError


class AtivacaoClientTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.path = Path(self.temporary.name) / "ATIVACAO.xlsx"
        AtivacaoClient._caminho_cache = None
        self.client = AtivacaoClient(self.path)

    def tearDown(self):
        AtivacaoClient._caminho_cache = None
        self.temporary.cleanup()

    @staticmethod
    def payload(**changes):
        values = {
            "cliente": "CLIENTE TESTE",
            "cidade": "CURITIBA",
            "servico": "IP DEDICADO",
            "tecnologia": "GPON",
            "empresa": "VIVO",
            "status": "OK",
            "data_agendamento": date.today().isoformat(),
            "data_execucao": date.today().isoformat(),
            "no_mes": "SIM",
            "tecnico": "TECNICO TESTE",
            "faturado": "SIM",
            "com_rfs": "NÃO",
            "situacao": "",
        }
        values.update(changes)
        return values

    def test_creates_expected_existing_workbook_schema(self):
        self.client.list()
        workbook = load_workbook(self.path)
        sheet = workbook["ATIVACAO"]
        self.assertEqual([sheet.cell(1, col).value for col in range(1, 15)], EXCEL_HEADERS)
        self.assertIn("ATIVACAO", sheet.tables)
        self.assertIn("CIDADES", workbook.sheetnames)
        self.assertEqual(workbook["CIDADES"]["A1"].value, "CIDADES")
        workbook.close()

    def test_create_three_edit_view_and_delete(self):
        first = self.client.create(self.payload(cliente="ALFA", situacao="FALTA CONFIG"))
        second = self.client.create(self.payload(cliente="BETA", empresa="RS TELECOM", status="NOK"))
        third = self.client.create(self.payload(cliente="GAMA", servico="SIP"))
        self.assertEqual([first["id"], second["id"], third["id"]], [1, 2, 3])
        self.assertEqual(self.client.list()["total"], 3)
        self.assertEqual(self.client.get(2)["cliente"], "BETA")
        self.assertEqual(self.client.get(1)["situacao"], "FALTA CONFIG")

        updated = self.client.update(2, {"status": "OK", "faturado": "NÃO", "situacao": "SEM ACESSO ERB"})
        self.assertEqual(updated["status"], "OK")
        self.assertEqual(updated["faturado"], "NÃO")
        self.assertEqual(updated["cidade"], "CURITIBA")
        self.assertEqual(updated["situacao"], "SEM ACESSO ERB")

        self.assertEqual(self.client.delete(1), {"id": 1, "deleted": True})
        self.assertEqual(self.client.list()["total"], 2)
        with self.assertRaises(DataClientError):
            self.client.get(1)

    def test_manual_id_can_be_created_and_changed_without_duplicates(self):
        created = self.client.create(self.payload(id="ATV-2224/A#01", cliente="ID MANUAL"))
        self.assertEqual(created["id"], "ATV-2224/A#01")

        with self.assertRaisesRegex(DataClientError, "Já existe uma atividade com o ID atv-2224/a#01"):
            self.client.create(self.payload(id="atv-2224/a#01", cliente="ID REPETIDO"))

        updated = self.client.update("ATV-2224/A#01", {"id": "OS 2026_01+X"})
        self.assertEqual(updated["id"], "OS 2026_01+X")
        with self.assertRaises(DataClientError):
            self.client.get("ATV-2224/A#01")

        leading_zero = self.client.create(self.payload(id="00123", cliente="ZERO À ESQUERDA"))
        self.assertEqual(leading_zero["id"], "00123")
        with self.assertRaisesRegex(DataClientError, "Campo ID obrigatório"):
            self.client.create(self.payload(id=""))

    def test_api_routes_accept_encoded_special_character_id(self):
        original_id = "OS/2026-A#01"
        self.client.create(self.payload(id=original_id, cliente="ROTA ESPECIAL"))
        app = create_app()
        app.config["TESTING"] = True

        with patch.object(api, "ativacao_client", self.client):
            http = app.test_client()
            encoded_id = quote(original_id, safe="")

            response = http.get(f"/api/ativacao/records/{encoded_id}")
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.get_json()["data"]["id"], original_id)

            changed_id = "RFS+ABC/02"
            response = http.patch(
                f"/api/ativacao/records/{encoded_id}",
                json={"id": changed_id},
            )
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.get_json()["data"]["id"], changed_id)

            response = http.delete(
                f"/api/ativacao/records/{quote(changed_id, safe='')}"
            )
            self.assertEqual(response.status_code, 200)

    def test_ativacao_page_exposes_situacao_in_table_and_both_modals(self):
        app = create_app()
        app.config["TESTING"] = True
        response = app.test_client().get("/ativacao")
        self.assertEqual(response.status_code, 200)
        html = response.get_data(as_text=True)
        self.assertIn("<th>Situação</th>", html)
        self.assertIn('id="newSituacao"', html)
        self.assertIn('id="editSituacao"', html)
        self.assertIn('id="fMesExecucao"', html)
        self.assertIn('id="newPermitirIdDuplicadoSim"', html)
        self.assertIn('id="editPermitirIdDuplicadoSim"', html)
        self.assertIn('id="editOriginalRow"', html)
        self.assertEqual(html.count("Aceita datas futuras."), 4)
        self.assertIn('<select id="newServico"', html)
        self.assertIn('<select id="editServico"', html)

    def test_duplicate_ids_can_be_managed_by_excel_row_when_authorized(self):
        first = self.client.create(self.payload(id="DUP-01", cliente="PRIMEIRO"))
        with self.assertRaisesRegex(DataClientError, "Já existe uma atividade"):
            self.client.create(self.payload(id="DUP-01", cliente="BLOQUEADO"))

        second = self.client.create(self.payload(
            id="DUP-01",
            cliente="SEGUNDO",
            permitir_id_duplicado=True,
        ))
        self.assertNotEqual(first["row_number"], second["row_number"])
        self.assertEqual(second["cliente"], "SEGUNDO")

        updated = self.client.update(
            "DUP-01",
            {"cliente": "SEGUNDO EDITADO"},
            row_number=second["row_number"],
        )
        self.assertEqual(updated["cliente"], "SEGUNDO EDITADO")
        self.assertEqual(self.client.get("DUP-01", first["row_number"])["cliente"], "PRIMEIRO")

        self.client.delete("DUP-01", row_number=first["row_number"])
        remaining = self.client.list({"q": "DUP-01"})["items"]
        self.assertEqual(len(remaining), 1)
        self.assertEqual(remaining[0]["cliente"], "SEGUNDO EDITADO")

    def test_future_dates_are_accepted_on_create_and_update(self):
        created = self.client.create(self.payload(
            id="FUTURO-01",
            data_agendamento="2099-12-30",
            data_execucao="2099-12-31",
        ))
        self.assertEqual(created["data_agendamento"], "2099-12-30")
        self.assertEqual(created["data_execucao"], "2099-12-31")

        updated = self.client.update(
            "FUTURO-01",
            {"data_agendamento": "2100-01-15", "data_execucao": "2100-01-20"},
            row_number=created["row_number"],
        )
        self.assertEqual(updated["data_agendamento"], "2100-01-15")
        self.assertEqual(updated["data_execucao"], "2100-01-20")

    def test_filters_by_execution_and_schedule_dates(self):
        today = date.today()
        yesterday = today - timedelta(days=1)
        self.client.create(self.payload(cliente="EXECUTA HOJE", data_agendamento=yesterday.isoformat()))
        self.client.create(self.payload(cliente="AGENDA HOJE", data_execucao=yesterday.isoformat()))

        execution = self.client.list({"data_execucao": today.isoformat()})
        schedule = self.client.list({"data_agendamento": today.isoformat()})
        self.assertEqual([item["cliente"] for item in execution["items"]], ["EXECUTA HOJE"])
        self.assertEqual([item["cliente"] for item in schedule["items"]], ["AGENDA HOJE"])

    def test_filters_by_execution_month_through_api(self):
        self.client.create(self.payload(id="MES-ATUAL", cliente="MÊS ATUAL", data_execucao="2026-09-15"))
        self.client.create(self.payload(id="MES-ANTERIOR", cliente="MÊS ANTERIOR", data_execucao="2026-08-31"))
        app = create_app()
        app.config["TESTING"] = True

        with patch.object(api, "ativacao_client", self.client):
            response = app.test_client().get("/api/ativacao/records?mesExecucao=2026-09")

        self.assertEqual(response.status_code, 200)
        data = response.get_json()["data"]
        self.assertEqual(data["total"], 1)
        self.assertEqual(data["items"][0]["id"], "MES-ATUAL")

        with self.assertRaisesRegex(DataClientError, "Mês da Execução inválido"):
            self.client.list({"mes_execucao": "2026-13"})

    def test_dashboard_requested_distributions(self):
        today = date.today().isoformat()
        self.client.create(self.payload(servico="IP DEDICADO", tecnologia="GPON", faturado="SIM", com_rfs="SIM"))
        self.client.create(self.payload(servico="SIP", tecnologia="ERB", status="NOK", faturado="NÃO", com_rfs="NÃO"))
        dashboard = self.client.dashboard({"data_execucao": today})
        self.assertEqual(dashboard["kpis"]["total"], 2)
        self.assertEqual(dashboard["por_status"]["values"], [1, 1, 0, 0])
        self.assertEqual(dashboard["por_tecnologia"]["values"], [1, 1])
        self.assertEqual(dashboard["por_faturado"]["values"], [1, 1])
        self.assertEqual(dashboard["por_rfs"]["values"], [1, 1])
        self.assertEqual(dashboard["por_data_execucao"], {"labels": [today], "values": [2]})
        self.assertEqual(set(dashboard["por_servico"]["labels"]), {"IP DEDICADO", "SIP"})

    def test_validation_and_dynamic_options(self):
        self.client.list()
        workbook = load_workbook(self.path)
        workbook["CIDADES"].append(["São José dos Pinhais"])
        workbook["CIDADES"].append(["Ponta Grossa"])
        workbook.save(self.path)
        workbook.close()
        self.client.create(self.payload(cliente="FUST", cidade="PINHAIS", servico="SIP", tecnico="EDSON"))
        options = self.client.options()
        self.assertIn("FUST", options["clientes"])
        self.assertIn("PINHAIS", options["cidades"])
        self.assertIn("SÃO JOSÉ DOS PINHAIS", options["cidades"])
        self.assertIn("PONTA GROSSA", options["cidades"])
        self.assertTrue(
            {"REPARO", "MIGRAÇÃO", "QUALIDADE", "SIP", "IP DEDICADO"}.issubset(options["servicos"])
        )
        self.assertEqual(
            options["tecnicos"],
            [
                "ALEXSANDRO NUNES DA SILVA",
                "ANDRE DOS SANTOS CARVALHO",
                "EDSON LUIS DE CRISTO",
                "EMERSON FAVERO RODRIGUES",
                "LUCAS SILVA ANDRADE",
                "MARCOS ROBERTO HOLTMAN",
                "ERENILSON SANT'ANA",
                "RS TELECOM",
                "STAFF",
                "BAIXADO REGIONAL",
            ],
        )
        self.assertEqual(options["status"], ["OK", "NOK", "FECHAMENTO INTERNO", "NÃO INSTALADO"])
        fechamento = self.client.create(
            self.payload(id="FECHAMENTO-01", status="FECHAMENTO INTERNO", tecnico="STAFF")
        )
        self.assertEqual(fechamento["status"], "FECHAMENTO INTERNO")
        self.assertEqual(fechamento["tecnico"], "STAFF")
        nao_instalado = self.client.create(self.payload(id="NAO-INSTALADO-01", status="NÃO INSTALADO"))
        self.assertEqual(nao_instalado["status"], "NÃO INSTALADO")
        self.assertEqual(
            options["situacoes"],
            [
                "FALTA GOLD JUMPER", "FALTA CONFIG", "SEM ACESSO",
                "FALTA VALIDAR", "EQPTO NÃO RETIRADO", "FALTA REDE",
                "LOCAL FECHADO", "LOCAL NÃO EXISTE", "AGENDA FUTURA",
                "VT VERSIONAMENTO", "SEM ACESSO ERB", "PENDÊNCIA CLIENTE",
            ],
        )
        with self.assertRaisesRegex(DataClientError, "Campo Cliente obrigatório"):
            self.client.create(self.payload(cliente=""))
        with self.assertRaisesRegex(DataClientError, "Data de Agendamento inválida"):
            self.client.create(self.payload(data_agendamento="31/02/2026"))
        with self.assertRaisesRegex(DataClientError, "Valor inválido para Situação"):
            self.client.create(self.payload(situacao="OUTRA"))

    def test_existing_workbook_receives_situacao_column_without_losing_rows(self):
        self.client.create(self.payload(id="LEG-01", cliente="CLIENTE ANTIGO"))
        workbook = load_workbook(self.path)
        sheet = workbook["ATIVACAO"]
        sheet.delete_cols(14, 1)
        sheet.tables["ATIVACAO"].ref = "A1:M2"
        sheet.auto_filter.ref = "A1:M2"
        workbook.save(self.path)
        workbook.close()

        result = self.client.list()
        self.assertEqual(result["items"][0]["cliente"], "CLIENTE ANTIGO")
        self.assertEqual(result["items"][0]["situacao"], "")

        workbook = load_workbook(self.path)
        sheet = workbook["ATIVACAO"]
        self.assertEqual(sheet["N1"].value, "SITUAÇÃO")
        self.assertEqual(sheet.tables["ATIVACAO"].ref, "A1:N2")
        workbook.close()

    def test_automatic_discovery_finds_file_on_desktop(self):
        desktop = Path(self.temporary.name) / "Desktop"
        desktop.mkdir()
        expected = desktop / "ATIVACAO.xlsx"
        expected.write_bytes(b"arquivo-portatil")

        with (
            patch.dict(
                os.environ,
                {"ATIVACAO_PATH": "", "ATIVACAO_FILENAME": "ATIVACAO.xlsx"},
                clear=False,
            ),
            patch.object(AtivacaoClient, "_candidate_directories", return_value=[desktop]),
            patch(
                "ativacao_client.DataClient._resolver_caminho_excel",
                side_effect=DataClientError("Planilha principal indisponível", 500),
            ),
            patch("ativacao_client._localizar_excel_no_onedrive", return_value=(None, [])),
        ):
            self.assertEqual(AtivacaoClient()._resolve_path(), expected)

    def test_automatic_discovery_rechecks_when_cached_file_was_moved(self):
        old_file = Path(self.temporary.name) / "antigo" / "ATIVACAO.xlsx"
        new_directory = Path(self.temporary.name) / "novo"
        new_directory.mkdir()
        new_file = new_directory / "ATIVACAO.xlsx"
        new_file.write_bytes(b"arquivo-movido")
        AtivacaoClient._caminho_cache = str(old_file)

        with (
            patch.dict(
                os.environ,
                {"ATIVACAO_PATH": "", "ATIVACAO_FILENAME": "ATIVACAO.xlsx"},
                clear=False,
            ),
            patch.object(
                AtivacaoClient, "_candidate_directories", return_value=[new_directory]
            ),
            patch(
                "ativacao_client.DataClient._resolver_caminho_excel",
                side_effect=DataClientError("Planilha principal indisponível", 500),
            ),
            patch("ativacao_client._localizar_excel_no_onedrive", return_value=(None, [])),
        ):
            self.assertEqual(AtivacaoClient()._resolve_path(), new_file)

    def test_automatic_mode_does_not_create_an_empty_file_when_missing(self):
        empty_directory = Path(self.temporary.name) / "vazio"
        empty_directory.mkdir()

        with (
            patch.dict(
                os.environ,
                {"ATIVACAO_PATH": "", "ATIVACAO_FILENAME": "ATIVACAO.xlsx"},
                clear=False,
            ),
            patch.object(
                AtivacaoClient, "_candidate_directories", return_value=[empty_directory]
            ),
            patch(
                "ativacao_client.DataClient._resolver_caminho_excel",
                side_effect=DataClientError("Planilha principal indisponível", 500),
            ),
            patch("ativacao_client._localizar_excel_no_onedrive", return_value=(None, [])),
        ):
            with self.assertRaisesRegex(DataClientError, "não foi encontrado"):
                AtivacaoClient()._ensure_workbook()
            self.assertFalse((empty_directory / "ATIVACAO.xlsx").exists())


if __name__ == "__main__":
    unittest.main()
