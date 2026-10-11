import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import api
from app import create_app
from backlog_client import BacklogClient, _number


HEADERS = [
    "Pedido", "ID_Cliente", "Cliente", "UF", "Cidade", "Carteira", "Servico",
    "Tecnologia_Report", "Dias_CarteiraAtual", "Data_Entrada", "Data_RFS", "Data_Atualizacao",
]


class BacklogTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name) / "Backlog.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Planilha1"
        sheet.append(HEADERS)
        sheet.append(["P-1", "C-1", "Cliente Um", "PR", "Curitiba", "Ativação", "IP", "GPON", 10, datetime(2026, 9, 1), datetime(2026, 9, 10), datetime(2026, 9, 30)])
        sheet.append(["P-2", "C-2", "Cliente Dois", "SP", "São Paulo", "PCC", "SIP", "ERB", 35, datetime(2026, 9, 2), "", datetime(2026, 9, 29)])
        sheet.append(["P-3", "C-3", "Cliente Três", "PR", "Pinhais", "Ativação", "IP", "GPON", 20, datetime(2026, 9, 3), datetime(2026, 9, 10), datetime(2026, 9, 30)])
        workbook.save(self.path)
        workbook.close()
        self.client = BacklogClient(self.path)
        self.patch = patch.object(api, "backlog_client", self.client)
        self.patch.start()
        app = create_app()
        app.config["TESTING"] = True
        self.http = app.test_client()

    def tearDown(self):
        self.patch.stop()
        self.temp.cleanup()

    def test_page_and_status(self):
        page = self.http.get("/backlog")
        self.assertEqual(page.status_code, 200)
        self.assertIn(b"Backlog", page.data)
        status = self.http.get("/api/backlog/status").get_json()["data"]
        self.assertEqual(status["arquivo"], "Backlog.xlsx")

    def test_dashboard_is_home_and_rede_has_its_own_route(self):
        home = self.http.get("/")
        rede = self.http.get("/rede")
        self.assertEqual(home.status_code, 200)
        self.assertIn(b"Dashboard B2B", home.data)
        self.assertEqual(rede.status_code, 200)
        self.assertIn(b"<h1>Rede</h1>", rede.data)

    def test_dashboard_aggregates_rfs_and_wallet_situation_by_quantity(self):
        data = self.client.dashboard()
        self.assertEqual(data["total"], 3)
        self.assertEqual(data["com_rfs"], 2)
        self.assertEqual(data["percentual_rfs"], 66.7)
        self.assertNotIn("aging_medio", data)
        self.assertEqual(data["rfs_por_dia"]["values"], [2])
        self.assertEqual(data["carteiras"][0]["nome"], "Ativação")
        self.assertEqual(data["carteiras"][0]["quantidade"], 2)
        self.assertEqual(data["carteiras"][0]["nivel"], "ok")

    def test_filters_pagination_and_export(self):
        response = self.http.get("/api/backlog/records?uf=PR&page=1&page_size=10")
        data = response.get_json()["data"]
        self.assertEqual(data["total"], 2)
        self.assertTrue(all(item["uf"] == "PR" for item in data["items"]))
        filtered = self.client.list_rows({"carteira": "PCC"})
        self.assertEqual(filtered["items"][0]["pedido"], "P-2")
        exported = self.http.get("/api/backlog/export?tecnologia=GPON")
        self.assertEqual(exported.status_code, 200)
        self.assertIn("spreadsheetml.sheet", exported.content_type)

        multiple = self.http.get("/api/backlog/records?uf=PR,SP&page=1&page_size=10").get_json()["data"]
        self.assertEqual(multiple["total"], 3)

    def test_wallet_situation_turns_critical_above_twenty_items(self):
        workbook = load_workbook(self.path)
        sheet = workbook["Planilha1"]
        for index in range(19):
            sheet.append([f"EXTRA-{index}", f"X-{index}", "Cliente", "PR", "Curitiba", "Ativação", "IP", "GPON", 1, datetime(2026, 9, 4), "", datetime(2026, 9, 30)])
        workbook.save(self.path)
        workbook.close()
        carteira = next(item for item in self.client.dashboard()["carteiras"] if item["nome"] == "Ativação")
        self.assertEqual(carteira["quantidade"], 21)
        self.assertEqual(carteira["nivel"], "critical")

    def test_options(self):
        data = self.http.get("/api/backlog/options").get_json()["data"]
        self.assertEqual(data["ufs"], ["PR", "SP"])
        self.assertEqual(data["tecnologias"], ["ERB", "GPON"])

    def test_numeric_values_keep_excel_decimals_and_accept_brazilian_text(self):
        self.assertEqual(_number(1.8), 1.8)
        self.assertEqual(_number("1,8"), 1.8)
        self.assertEqual(_number("1.234,5"), 1234.5)


if __name__ == "__main__":
    unittest.main()
