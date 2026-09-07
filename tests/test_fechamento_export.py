# -*- coding: utf-8 -*-
"""Testes da exportação formatada do quadro Fechamento Geral."""

import hashlib
import os
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from excel_client import DataClient, FIELDS


class FechamentoExportTests(unittest.TestCase):
    def _criar_base(self, path: Path):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "REDEB2B"
        headers = [*FIELDS, "COLUNA_EXTRA"]
        worksheet.append(headers)

        def linha(item_id, atividade, status, agendamento, conclusao, extra):
            values = {field: "" for field in FIELDS}
            values.update({
                "IDCLIENTE": item_id,
                "CLIENTE": f"Cliente {item_id}",
                "ATIVIDADE": atividade,
                "STATUS": status,
                "DATAAGENDAMENTO": agendamento,
                "DATACONCLUSAO": conclusao,
            })
            return [values[field] for field in FIELDS] + [extra]

        # A primeira linha é excluída. As duas seguintes precisam subir de
        # posição na cópia, exercitando também a tradução das fórmulas.
        worksheet.append(linha("0", "REPARO", "NOVO", datetime(2026, 8, 24), "", "=LEN(B2)"))
        worksheet.append(linha("1", "REPARO", "CONCLUÍDO", datetime(2026, 8, 1), datetime(2026, 8, 24), "=LEN(B3)"))
        worksheet.append(linha("2", "ESTEIRA", "PCC", datetime(2026, 8, 24), "", "=LEN(B4)"))
        worksheet.append(linha("3", "INSTALAÇÃO", "AGENDADO", datetime(2026, 8, 24), "", "=LEN(B5)"))
        worksheet.append(linha("4", "REPARO", "CONCLUIDO", datetime(2026, 8, 24), datetime(2026, 8, 23), "=LEN(B6)"))

        table = Table(displayName="TabelaRede", ref=f"A1:V{worksheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium4",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)
        worksheet.freeze_panes = "A2"
        worksheet.column_dimensions["B"].width = 31
        worksheet["A3"].fill = PatternFill("solid", fgColor="FFF2CC")
        workbook.save(path)
        workbook.close()

    @staticmethod
    def _digest(path: Path):
        return hashlib.sha256(path.read_bytes()).hexdigest()

    def test_export_preserva_tabela_e_contem_apenas_registros_do_fechamento(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            source = Path(tmp_dir) / "REDE_B2B.xlsx"
            self._criar_base(source)
            digest_before = self._digest(source)

            env = {
                "EXCEL_PATH": str(source),
                "EXCEL_SHEET_NAME": "REDEB2B",
                "EXCEL_TABLE": "TabelaRede",
            }
            with patch.dict(os.environ, env, clear=False):
                client = DataClient()
                buffer, metadata = client.export_fechamento_geral(
                    data_inicio="2026-08-24",
                    data_fim="2026-08-24",
                )

            self.assertEqual(digest_before, self._digest(source), "A base original foi alterada")
            self.assertEqual(metadata["total"], 2)

            exported = load_workbook(buffer, data_only=False)
            try:
                worksheet = exported["REDEB2B"]
                table = worksheet.tables["TabelaRede"]
                self.assertEqual(table.ref, "A1:V3")
                self.assertEqual(table.tableStyleInfo.name, "TableStyleMedium4")
                self.assertTrue(table.tableStyleInfo.showRowStripes)
                self.assertEqual(worksheet.freeze_panes, "A2")
                self.assertEqual(worksheet.column_dimensions["B"].width, 31)

                ids = [worksheet.cell(row, 1).value for row in range(2, 4)]
                self.assertEqual(ids, ["1", "2"])
                self.assertEqual(worksheet["V2"].value, "=LEN(B2)")
                self.assertEqual(worksheet["V3"].value, "=LEN(B3)")
                self.assertEqual(worksheet["A2"].fill.fgColor.rgb, "00FFF2CC")
            finally:
                exported.close()

    def test_export_sem_registros_mantem_tabela_valida_com_linha_vazia(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            source = Path(tmp_dir) / "REDE_B2B.xlsx"
            self._criar_base(source)
            digest_before = self._digest(source)

            env = {
                "EXCEL_PATH": str(source),
                "EXCEL_SHEET_NAME": "REDEB2B",
                "EXCEL_TABLE": "TabelaRede",
            }
            with patch.dict(os.environ, env, clear=False):
                client = DataClient()
                buffer, metadata = client.export_fechamento_geral(
                    data_inicio="2026-09-01",
                    data_fim="2026-09-01",
                )

            self.assertEqual(digest_before, self._digest(source), "A base original foi alterada")
            self.assertEqual(metadata["total"], 0)

            exported = load_workbook(buffer, data_only=False)
            try:
                worksheet = exported["REDEB2B"]
                table = worksheet.tables["TabelaRede"]
                self.assertEqual(table.ref, "A1:V2")
                self.assertEqual(table.tableStyleInfo.name, "TableStyleMedium4")
                self.assertTrue(all(worksheet.cell(2, col).value is None for col in range(1, 23)))
            finally:
                exported.close()


if __name__ == "__main__":
    unittest.main()
