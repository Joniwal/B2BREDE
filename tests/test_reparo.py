import tempfile
import unittest
import zipfile
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import patch
from urllib.parse import quote

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

import api
from app import create_app
from reparo_client import EXCEL_HEADERS, ReparoClient


class ReparoClientTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.path = Path(self.temporary.name) / "REPARO.xlsx"
        ReparoClient._caminho_cache = None
        self.client = ReparoClient(self.path)

    def tearDown(self):
        ReparoClient._caminho_cache = None
        self.temporary.cleanup()

    @staticmethod
    def payload(**changes):
        data = {
            "BD": "BD-1001",
            "ID_VANTIVE": "V-1001",
            "STATUS": "Aberto",
            "CLIENTE": "CLIENTE TESTE",
            "CIDADE": "CURITIBA",
            "UF": "PR",
            "DATA_ABERTURA": "2026-09-24",
            "MOTIVO_REAL_BD": "ROMPIMENTO DE FIBRA",
            "TMR": "2.5",
            "EPS": "VIVO",
            "TECNOLOGIA": "GPON",
            "TECNICO": "TÉCNICO TESTE",
            "REINC_30D": "Não",
        }
        data.update(changes)
        return data

    def test_creates_shared_excel_base_and_crud(self):
        created = self.client.create(self.payload())
        self.assertEqual(created["ID_VANTIVE"], "V-1001")
        self.assertEqual(created["DATA_ABERTURA"], "2026-09-24")

        workbook = load_workbook(self.path)
        sheet = workbook["REPARO"]
        self.assertEqual([sheet.cell(1, index).value for index in range(1, len(EXCEL_HEADERS) + 1)], EXCEL_HEADERS)
        self.assertEqual(len(sheet.tables), 0)
        self.assertEqual(sheet.auto_filter.ref, f"A1:{sheet.cell(1, len(EXCEL_HEADERS)).column_letter}2")
        self.assertIn("LISTAS", workbook.sheetnames)
        workbook.close()

        updated = self.client.update(
            "V-1001",
            {"STATUS": "Encerrado", "DATA_ENCERRAMENTO": "2026-09-25"},
            row_number=created["row_number"],
        )
        self.assertEqual(updated["STATUS"], "Encerrado")
        self.assertEqual(updated["DATA_ENCERRAMENTO"], "2026-09-25")
        self.assertEqual(updated["TMR_PARADA"], "24.00")
        self.assertTrue(updated["ATUALIZADO_EM"])

        self.assertEqual(self.client.list({"status": "Encerrado"})["total"], 1)
        dashboard = self.client.dashboard()
        self.assertEqual(dashboard["kpis"]["total"], 1)
        self.assertEqual(dashboard["kpis"]["encerrados"], 1)
        self.assertEqual(dashboard["kpis"]["tmr_medio"], 2.5)

        self.client.delete("V-1001", row_number=created["row_number"])
        self.assertEqual(self.client.list()["total"], 0)

    def test_reparo_form_fields_and_tmr_parada_calculation(self):
        created = self.client.create(self.payload(
            ID_VANTIVE="R-ERB-01",
            DATA_ABERTURA="2026-09-24T08:00",
            DATA_ENCERRAMENTO="2026-09-24T18:00",
            TIPO_DEFEITO="CAMPO",
            PARADA_RELOGIO="SIM",
            INICIO="2026-09-24T12:00",
            FIM="2026-09-24T13:00",
            TIPO_ARD_ERB="ERB",
            ARD_ERB="CTA123",
            USUARIO_BAIXA="OPERADOR",
        ))
        self.assertEqual(created["TMR_PARADA"], "9.00")
        self.assertEqual(created["TIPO_ARD_ERB"], "ERB")
        self.assertEqual(created["ARD_ERB"], "CTA123")
        self.assertTrue(created["ATUALIZADO_EM"])

        with self.assertRaisesRegex(Exception, "máximo 10 caracteres"):
            self.client.update(
                "R-ERB-01", {"ARD_ERB": "SIGLA-MUITO-LONGA"},
                row_number=created["row_number"],
            )

        javascript = (Path(__file__).parents[1] / "static" / "js" / "reparo.js").read_text(encoding="utf-8")
        self.assertIn('select("EPS","EPS",EPS_OPTIONS)', javascript)
        self.assertIn('type="radio" name="tipoArdErb"', javascript)
        self.assertIn('maxlength="10"', javascript)
        self.assertIn('field("TMR-(PARADA) em horas","TMR_PARADA"', javascript)

    def test_form_options_are_read_from_listas_sheet(self):
        self.client.create(self.payload(ID_VANTIVE="R-LISTAS-01"))
        workbook = load_workbook(self.path)
        sheet = workbook["LISTAS"]
        headers = {
            sheet.cell(1, column).value: column
            for column in range(1, sheet.max_column + 1)
        }
        values = {
            "MOTIVO REAL DO BD": "ROMPIMENTO",
            "TÉCNICO": "TECNICO A",
            "FILA DE ENCERRAMENTO": "FILA N1",
            "TECNOLOGIA": "GPON",
            "USUÁRIO": "OPERADOR A",
        }
        for header, value in values.items():
            sheet.cell(2, headers[header], value)
        workbook.save(self.path)
        workbook.close()

        options = self.client.options()
        self.assertEqual(options["motivo_real_bd"], ["ROMPIMENTO"])
        self.assertEqual(options["tecnico"], ["TECNICO A"])
        self.assertEqual(options["fila_encerramento"], ["FILA N1"])
        self.assertEqual(options["tecnologia"], ["GPON"])
        self.assertEqual(options["usuario"], ["OPERADOR A"])

    def test_dashboard_groups_repeated_lp_15_from_last_30_days(self):
        today = date.today()
        common = {
            "LP_15": "LP-7788",
            "CLIENTE": "CLIENTE RECORRENTE",
        }
        self.client.create(self.payload(
            ID_VANTIVE="R-LP-01", DATA_ABERTURA=today.isoformat(),
            MOTIVO_REAL_BD="ROMPIMENTO", **common,
        ))
        self.client.create(self.payload(
            ID_VANTIVE="R-LP-02", DATA_ABERTURA=(today - timedelta(days=5)).isoformat(),
            MOTIVO_REAL_BD="ROMPIMENTO", **common,
        ))
        self.client.create(self.payload(
            ID_VANTIVE="R-LP-03", DATA_ABERTURA=(today - timedelta(days=12)).isoformat(),
            MOTIVO_REAL_BD="MASSIVA", **common,
        ))
        self.client.create(self.payload(
            ID_VANTIVE="R-LP-ANTIGO", DATA_ABERTURA=(today - timedelta(days=40)).isoformat(),
            MOTIVO_REAL_BD="ANTIGO", **common,
        ))
        self.client.create(self.payload(
            ID_VANTIVE="R-LP-UNICO", DATA_ABERTURA=today.isoformat(),
            LP_15="LP-UNICA", CLIENTE="CLIENTE ÚNICO",
        ))

        analysis = self.client.dashboard()["reincidencias_lp_30d"]
        self.assertEqual(analysis["data_inicio"], (today - timedelta(days=29)).isoformat())
        self.assertEqual(analysis["data_fim"], today.isoformat())
        self.assertEqual(len(analysis["itens"]), 1)
        repeated = analysis["itens"][0]
        self.assertEqual(repeated["lp_15"], "LP-7788")
        self.assertEqual(repeated["cliente"], "CLIENTE RECORRENTE")
        self.assertEqual(repeated["quantidade"], 3)
        self.assertEqual(repeated["motivos"], [
            {"motivo": "ROMPIMENTO", "quantidade": 2},
            {"motivo": "MASSIVA", "quantidade": 1},
        ])

    def test_existing_excel_base_is_prepared_without_losing_rows(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "BD"
        headers = ["BD", "ID_VANTIVE", "STATUS", "CLIENTE", "DATA_ABERTURA", "MOTIVO_REAL_BD"]
        sheet.append(headers)
        sheet.append(["BD-9", "V-9", "Aberto", "CLIENTE EXISTENTE", "2026-09-20", "FALHA"])
        workbook.save(self.path)
        workbook.close()

        listed = self.client.list()
        self.assertEqual(listed["total"], 1)
        self.assertEqual(listed["items"][0]["CLIENTE"], "CLIENTE EXISTENTE")

        prepared = load_workbook(self.path)
        sheet = prepared["BD"]
        self.assertEqual(len(sheet.tables), 0)
        self.assertEqual(sheet.auto_filter.ref, f"A1:{sheet.cell(1, sheet.max_column).column_letter}2")
        normalized_headers = [sheet.cell(1, index).value for index in range(1, sheet.max_column + 1)]
        self.assertIn("TECNOLOGIA", normalized_headers)
        self.assertEqual(sheet["D2"].value, "CLIENTE EXISTENTE")
        prepared.close()

    def test_replaces_structured_table_with_excel_compatible_filter(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Reparo"
        sheet.append(EXCEL_HEADERS)
        row = [""] * len(EXCEL_HEADERS)
        row[EXCEL_HEADERS.index("ID_VANTIVE")] = "SEM-REPARO-001"
        row[EXCEL_HEADERS.index("CLIENTE")] = "CLIENTE PRESERVADO"
        sheet.append(row)
        last_column = sheet.cell(1, len(EXCEL_HEADERS)).column_letter
        table = Table(displayName="REPARO", ref=f"A1:{last_column}2")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium4", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=False,
        )
        sheet.add_table(table)
        workbook.save(self.path)
        workbook.close()

        listed = self.client.list()
        self.assertEqual(listed["total"], 1)
        self.assertEqual(listed["items"][0]["CLIENTE"], "CLIENTE PRESERVADO")

        repaired = load_workbook(self.path)
        sheet = repaired["Reparo"]
        self.assertEqual(len(sheet.tables), 0)
        self.assertEqual(sheet.auto_filter.ref, f"A1:{last_column}2")
        self.assertEqual(sheet.cell(2, EXCEL_HEADERS.index("CLIENTE") + 1).value, "CLIENTE PRESERVADO")
        repaired.close()
        with zipfile.ZipFile(self.path) as archive:
            self.assertFalse(any(name.startswith("xl/tables/") for name in archive.namelist()))

    def test_api_page_navigation_export_and_crud(self):
        app = create_app()
        app.config["TESTING"] = True
        http = app.test_client()

        page = http.get("/reparo")
        self.assertEqual(page.status_code, 200)
        html = page.get_data(as_text=True)
        self.assertIn("REPARO.xlsx", html)
        self.assertIn('href="/ativacao"', html)
        self.assertNotIn('data-view="importar"', html)
        self.assertNotIn("xlsx.full.min.js", html)
        self.assertNotIn('id="btnAtualizarReparo"', html)

        javascript = (Path(__file__).parents[1] / "static" / "js" / "reparo.js").read_text(encoding="utf-8")
        self.assertGreaterEqual(javascript.count('id="refreshData"'), 2)
        self.assertIn("Reparos repetidos por LP_15", javascript)
        self.assertIn(
            'const TABLE_COLS = ["BD","CLIENTE","STATUS","CIDADE","DATA_ABERTURA",'
            '"MOTIVO_REAL_BD","TMR","EPS","TECNOLOGIA"]',
            javascript,
        )
        self.assertNotIn('data-view-row=', javascript)
        self.assertIn('aria-label="Editar registro"', javascript)
        self.assertIn('aria-label="Excluir registro"', javascript)
        self.assertIn("renderNumberedPagination", javascript)
        self.assertNotIn('id="prevPg"', javascript)

        for route in ("/", "/dashboard", "/ativacao"):
            response = http.get(route)
            self.assertEqual(response.status_code, 200)
            self.assertIn('href="/reparo"', response.get_data(as_text=True))

        with patch.object(api, "reparo_client", self.client):
            created = http.post("/api/reparo/records", json=self.payload(ID_VANTIVE="V/ABC-02"))
            self.assertEqual(created.status_code, 201)
            item = created.get_json()["data"]

            listed = http.get("/api/reparo/records?status=Aberto")
            self.assertEqual(listed.status_code, 200)
            self.assertEqual(listed.get_json()["data"]["total"], 1)

            encoded = quote("V/ABC-02", safe="")
            changed = http.patch(
                f"/api/reparo/records/{encoded}?row={item['row_number']}",
                json={"STATUS": "Em Andamento"},
            )
            self.assertEqual(changed.status_code, 200)
            self.assertEqual(changed.get_json()["data"]["STATUS"], "Em Andamento")

            exported = http.get("/api/reparo/export")
            self.assertEqual(exported.status_code, 200)
            self.assertIn(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                exported.content_type,
            )

            imported = http.post("/api/reparo/import", json={"records": []})
            self.assertEqual(imported.status_code, 404)

            deleted = http.delete(f"/api/reparo/records/{encoded}?row={item['row_number']}")
            self.assertEqual(deleted.status_code, 200)


if __name__ == "__main__":
    unittest.main()
