import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from app import create_app
from excel_client import DataClient, FIELDS


class SiteLiberadoTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.path = Path(self.temporary.name) / "REDE_B2B.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "REDEB2B"
        headers = [
            "SITE" if field == "SITELIBERADO" else field
            for field in FIELDS
        ]
        sheet.append(headers)
        row = {field: "" for field in FIELDS}
        row.update({
            "IDCLIENTE": "ERB-001",
            "CLIENTE": "CLIENTE ERB",
            "TECNOLOGIA": "ERB",
            "SITELIBERADO": "SIM",
        })
        sheet.append([
            row["SITELIBERADO"] if header == "SITE" else row[header]
            for header in headers
        ])
        table = Table(
            displayName="REDEB2B",
            ref=f"A1:{get_column_letter(len(headers))}2",
        )
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
        workbook.save(self.path)
        workbook.close()

    def tearDown(self):
        DataClient._caminho_cache = None
        self.temporary.cleanup()

    def client(self):
        environment = {
            "EXCEL_PATH": str(self.path),
            "EXCEL_SHEET_NAME": "REDEB2B",
            "EXCEL_TABLE": "REDEB2B",
        }
        return patch.dict(os.environ, environment, clear=False)

    def test_site_liberado_alias_is_read_and_saved(self):
        with self.client():
            client = DataClient()
            self.assertEqual(client.get_item("ERB-001")["SITELIBERADO"], "SIM")

            updated = client.update_item(
                "ERB-001",
                {"TECNOLOGIA": "ERB", "SITELIBERADO": "não"},
            )
            self.assertEqual(updated["SITELIBERADO"], "NÃO")

            updated = client.update_item(
                "ERB-001",
                {"TECNOLOGIA": "ERB", "SITELIBERADO": "SIM"},
            )
            self.assertEqual(updated["SITELIBERADO"], "SIM")

            workbook = load_workbook(self.path)
            sheet = workbook["REDEB2B"]
            headers = {
                sheet.cell(1, column).value: column
                for column in range(1, sheet.max_column + 1)
            }
            self.assertEqual(sheet.cell(2, headers["SITE"]).value, "SIM")
            workbook.close()

            client.create_item({
                "IDCLIENTE": "GPON-002",
                "CLIENTE": "CLIENTE GPON",
                "TECNOLOGIA": "GPON",
            })
            filtered = client.list_items(filters={"tecnologia": "ERB"})
            self.assertEqual(filtered["total"], 1)
            self.assertEqual(filtered["items"][0]["IDCLIENTE"], "ERB-001")

            updated = client.update_item(
                "ERB-001",
                {"TECNOLOGIA": "GPON"},
            )
            self.assertEqual(updated["SITELIBERADO"], "")

        workbook = load_workbook(self.path)
        sheet = workbook["REDEB2B"]
        headers = {sheet.cell(1, column).value: column for column in range(1, sheet.max_column + 1)}
        self.assertIsNone(sheet.cell(2, headers["SITE"]).value)
        workbook.close()

    def test_page_contains_conditional_site_selector_and_technology_filter(self):
        app = create_app()
        app.config["TESTING"] = True
        response = app.test_client().get("/")
        self.assertEqual(response.status_code, 200)
        html = response.get_data(as_text=True)
        self.assertIn('id="siteLiberadoField"', html)
        self.assertIn('id="f_SITELIBERADO"', html)
        self.assertIn('<option value="SIM">SIM</option>', html)
        self.assertIn('id="fTecnologia"', html)
        self.assertIn(">Site</th>", html)

        javascript = (
            Path(__file__).parents[1] / "static" / "js" / "main.js"
        ).read_text(encoding="utf-8")
        self.assertIn("renderSiteLiberadoIcon", javascript)
        self.assertIn("bi-broadcast-pin", javascript)
        self.assertIn('tecnologia: state.filters.tecnologia', javascript)


if __name__ == "__main__":
    unittest.main()
